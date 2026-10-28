"""Bounded local spreadsheet formulas with independently computed OOXML caches."""

from __future__ import annotations

from io import BytesIO
import math
from pathlib import Path
import re
from typing import Any
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def cache_formulas(path: Path) -> None:
    book = load_workbook(path, data_only=False)
    values: dict[tuple[str, str], float] = {}
    active: set[tuple[str, str]] = set()

    def number(value: Any) -> float:
        if value is None:
            return 0.0
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ValueError("formula_requires_numeric_cells")
        return float(value)

    def cell_value(sheet: Any, coordinate: str) -> Any:
        cell = sheet[coordinate]
        if cell.data_type != "f":
            return cell.value
        key = (sheet.title, coordinate)
        if key in active or len(active) >= 50:
            raise ValueError("circular_or_deep_formula")
        if key not in values:
            active.add(key)
            values[key] = evaluate(sheet, cell.value)
            active.remove(key)
        return values[key]

    def reference(sheet: Any, text: str) -> list[Any]:
        if "!" in text:
            title, text = text.rsplit("!", 1)
            title = title.strip("'").replace("''", "'")
            if title not in book.sheetnames:
                raise ValueError("unknown_formula_sheet")
            sheet = book[title]
        if not re.fullmatch(
            r"\$?[A-Z]{1,3}\$?[1-9][0-9]{0,4}(?::\$?[A-Z]{1,3}\$?[1-9][0-9]{0,4})?",
            text,
        ):
            raise ValueError("unsupported_formula_reference")
        left, top, right, bottom = range_boundaries(text)
        if left is None or top is None or right is None or bottom is None:
            raise ValueError("unsupported_formula_reference")
        if (
            right > sheet.max_column
            or bottom > sheet.max_row
            or (right - left + 1) * (bottom - top + 1) > 10000
        ):
            raise ValueError("formula_range_limit")
        return [
            cell_value(sheet, cell.coordinate)
            for row in sheet.iter_rows(
                min_row=top, max_row=bottom, min_col=left, max_col=right
            )
            for cell in row
        ]

    def evaluate(sheet: Any, formula: str) -> float:
        tokens = [t for t in Tokenizer(formula).items if t.type != "WHITE-SPACE"]
        if len(tokens) > 200:
            raise ValueError("formula_size_limit")
        position = 0

        def expression() -> list[Any]:
            nonlocal position
            if position >= len(tokens):
                raise ValueError("incomplete_formula")
            token = tokens[position]
            position += 1
            if token.type == "OPERAND":
                if token.subtype == "NUMBER":
                    return [float(token.value)]
                if token.subtype == "TEXT":
                    return [token.value[1:-1].replace('""', '"')]
                if token.subtype == "RANGE":
                    return reference(sheet, token.value)
            if token.type == "FUNC" and token.subtype == "OPEN":
                name = token.value[:-1].upper()
                args = [expression()]
                while position < len(tokens) and tokens[position].type == "SEP":
                    position += 1
                    args.append(expression())
                if (
                    position >= len(tokens)
                    or tokens[position].type != "FUNC"
                    or tokens[position].subtype != "CLOSE"
                ):
                    raise ValueError("unsupported_formula")
                position += 1
                if name == "SUM":
                    return [
                        sum(
                            number(v)
                            for arg in args
                            for v in arg
                            if isinstance(v, (int, float)) and not isinstance(v, bool)
                        )
                    ]
                if name == "SUMIF" and len(args) in (2, 3) and len(args[1]) == 1:
                    criteria, criterion = args[0], args[1][0]
                    amounts = args[2] if len(args) == 3 else criteria
                    if len(criteria) != len(amounts) or (
                        isinstance(criterion, str)
                        and any(c in criterion for c in "*?~<>=!")
                    ):
                        raise ValueError("unsupported_sumif_criteria")
                    return [
                        sum(
                            number(v)
                            for c, v in zip(criteria, amounts, strict=True)
                            if str(c).casefold() == str(criterion).casefold()
                        )
                    ]
                raise ValueError("unsupported_formula_function")
            raise ValueError("unsupported_formula")

        result = expression()
        if position != len(tokens) or len(result) != 1:
            raise ValueError("unsupported_formula")
        value = number(result[0])
        if not math.isfinite(value):
            raise ValueError("nonfinite_formula")
        return value

    try:
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        cell_value(sheet, cell.coordinate)
        output = BytesIO()
        with ZipFile(path) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
            for name in source.namelist():
                raw = source.read(name)
                match = re.fullmatch(r"xl/worksheets/sheet([0-9]+)\.xml", name)
                if match:
                    sheet = book.worksheets[int(match[1]) - 1]
                    root = ET.fromstring(raw)
                    for xml_cell in root.iter(NS + "c"):
                        key = (sheet.title, xml_cell.attrib["r"])
                        if key in values:
                            xml_value = xml_cell.find(NS + "v")
                            if xml_value is None:
                                xml_value = ET.SubElement(xml_cell, NS + "v")
                            xml_value.text = repr(values[key])
                    raw = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(name, raw)
        path.write_bytes(output.getvalue())
        cached = load_workbook(path, data_only=True)
        try:
            for (sheet_name, coordinate), value in values.items():
                if cached[sheet_name][coordinate].value != value:
                    raise ValueError("formula_cache_verification_failed")
        finally:
            cached.close()
    finally:
        book.close()
