"""Bind table summaries to written data and check publication against source rows."""

from collections import Counter
from decimal import Decimal, InvalidOperation
from io import BytesIO
import math
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from .schema import Content

SUMIF = re.compile(
    r'^=SUMIF\((?:\'?[0-9]+\'?!)?\$?([A-Z]+)\$?\d+:\$?\1\$?\d+,\s*"([^"]+)",\s*(?:\'?[0-9]+\'?!)?\$?([A-Z]+)\$?\d+:\$?\3\$?\d+\)$',
    re.I,
)
SUM = re.compile(r"^=SUM\((?:'?[0-9]+'?!)?\$?([A-Z]+)\$?(\d+):\$?\1\$?(\d+)\)$", re.I)


def bind_summaries(content: Content) -> Content:
    value = content.model_dump()
    for index, section in enumerate(value["sections"]):
        written = (
            (1 if index == 0 else 0)
            + 1
            + len(section["paragraphs"])
            + len(section["bullets"])
        )
        for table in section["tables"]:
            rows = table["rows"]
            start = written + 2
            data = []
            for row in rows:
                if any(cell.startswith("=") for cell in row):
                    break
                data.append(row)
            end = start + len(data) - 1
            if table["formulas"] and data:
                replacements: dict[str, str] = {}
                for row in rows[len(data) :]:
                    for cell in row:
                        match = SUMIF.fullmatch(cell)
                        if (
                            match
                            and "!" not in cell
                            and max(
                                column_index_from_string(match[1]),
                                column_index_from_string(match[3]),
                            )
                            <= len(table["headers"])
                        ):
                            replacements[cell] = (
                                f'=SUMIF({match[1]}{start}:{match[1]}{end},"{match[2]}",{match[3]}{start}:{match[3]}{end})'
                            )
                # A range with exactly the data-row count is a table aggregate.
                # SUM over summary cells or explicit subsets remains unchanged.
                for row in rows[len(data) :]:
                    for cell in row:
                        match = SUM.fullmatch(cell)
                        if (
                            match
                            and "!" not in cell
                            and int(match[3]) - int(match[2]) + 1 == len(data)
                            and int(match[2]) < start + len(data)
                        ):
                            replacements[cell] = (
                                f"=SUM({match[1]}{start}:{match[1]}{end})"
                            )
                table["rows"] = tuple(
                    tuple(replacements.get(cell, cell) for cell in row) for row in rows
                )
            written += 1 + len(rows)
    return Content.model_validate(value)


def scalar(value: Any) -> str:
    text = str(value if value is not None else "").strip()
    try:
        number = Decimal(text)
        if number.is_finite():
            return str(number.normalize())
    except InvalidOperation:
        pass
    return text


def validate_sources(raw: bytes, selections: list[list[dict[str, Any]]]) -> None:
    """No source row may disappear; cached category totals use source amounts."""
    if not selections:
        return
    source_rows: list[tuple[str, ...]] = []
    for segments in selections:
        rows: dict[tuple[str, int], dict[int, str]] = {}
        for segment in segments:
            locator = segment["locator"]
            match = re.fullmatch(r"([A-Z]+)([0-9]+)", locator.get("cell", ""))
            if match:
                rows.setdefault((locator["sheet"], int(match[2])), {})[
                    column_index_from_string(match[1])
                ] = scalar(segment["text"])
        for cells in rows.values():
            source_row = tuple(
                cells.get(col, "") for col in range(min(cells), max(cells) + 1)
            )
            if any(
                re.fullmatch(r"-?\d+(?:\.\d+)?(?:E[+-]?\d+)?", cell)
                for cell in source_row
            ):
                source_rows.append(source_row)
    if not source_rows:
        raise ValueError("spreadsheet_source_mismatch:no_numeric_source_rows")
    formulas = load_workbook(BytesIO(raw), data_only=False)
    cached = load_workbook(BytesIO(raw), data_only=True)
    try:
        actual = Counter(
            tuple(scalar(cell.value) for cell in row)
            for sheet in formulas
            for row in sheet
            if not any(cell.data_type == "f" for cell in row)
        )
        required = Counter(source_rows)
        if required - actual:
            raise ValueError("spreadsheet_source_mismatch:missing_rows")
        for sheet in formulas:
            for row in sheet:
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    match = SUMIF.fullmatch(str(cell.value))
                    if not match:
                        if "SUMIF(" in str(cell.value).upper():
                            raise ValueError(
                                "spreadsheet_source_mismatch:unverifiable_formula"
                            )
                        total = SUM.fullmatch(str(cell.value))
                        if total:
                            col = column_index_from_string(total[1]) - 1
                            if all(col < len(source) for source in source_rows):
                                try:
                                    expected_total = sum(
                                        Decimal(source[col]) for source in source_rows
                                    )
                                except InvalidOperation:
                                    continue
                                actual_total = cached[sheet.title][
                                    cell.coordinate
                                ].value
                                if not isinstance(
                                    actual_total, (int, float)
                                ) or not math.isclose(
                                    float(expected_total),
                                    actual_total,
                                    rel_tol=1e-12,
                                    abs_tol=1e-12,
                                ):
                                    raise ValueError(
                                        "spreadsheet_source_mismatch:cached_total"
                                    )
                        continue
                    category, amount = (
                        column_index_from_string(match[1]) - 1,
                        column_index_from_string(match[3]) - 1,
                    )
                    if any(
                        max(category, amount) >= len(source) for source in source_rows
                    ):
                        raise ValueError("spreadsheet_source_mismatch:columns")
                    expected = sum(
                        Decimal(source[amount])
                        for source in source_rows
                        if source[category].casefold() == match[2].casefold()
                    )
                    actual_value = cached[sheet.title][cell.coordinate].value
                    if not isinstance(actual_value, (int, float)) or not math.isclose(
                        float(expected), actual_value, rel_tol=1e-12, abs_tol=1e-12
                    ):
                        raise ValueError("spreadsheet_source_mismatch:cached_total")
    finally:
        formulas.close()
        cached.close()
