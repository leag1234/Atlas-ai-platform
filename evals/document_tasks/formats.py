"""Deterministic synthetic formats and bounded archive I/O for evaluator fixtures."""

from __future__ import annotations

import csv
import gzip
import hashlib
import io
import json
import tarfile
from pathlib import Path, PurePosixPath
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.pdfgen.canvas import Canvas

from packages.doc_render.renderer import normalize_zip


def digest(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def json_bytes(value: Any) -> bytes:
    return (
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    ).encode()


def safe_path(name: str) -> str:
    if (
        not name
        or "\\" in name
        or "\0" in name
        or name.startswith("/")
        or any(p in {"", ".", ".."} for p in name.split("/"))
    ):
        raise ValueError("unsafe_relative_path")
    return str(PurePosixPath(name))


def pack(files: dict[str, bytes]) -> bytes:
    stream = io.BytesIO()
    with tarfile.open(fileobj=stream, mode="w", format=tarfile.USTAR_FORMAT) as archive:
        for name, raw in sorted(files.items()):
            info = tarfile.TarInfo(safe_path(name))
            info.size, info.mtime, info.mode = len(raw), 0, 0o600
            info.uid = info.gid = 0
            info.uname = info.gname = ""
            archive.addfile(info, io.BytesIO(raw))
    return gzip.compress(stream.getvalue(), compresslevel=9, mtime=0)


def unpack(raw: bytes, destination: Path, maximum: int = 32 * 1024 * 1024) -> None:
    """Preflight the entire outer archive; nested adversarial inputs stay opaque."""
    with tarfile.open(fileobj=io.BytesIO(raw), mode="r:gz") as archive:
        members = archive.getmembers()
        names = [safe_path(m.name) for m in members]
        if (
            len(members) > 2000
            or len(names) != len(set(names))
            or any(not m.isfile() or m.size < 0 for m in members)
            or sum(m.size for m in members) > maximum
        ):
            raise ValueError("unsafe_fixture_archive")
        for member in members:
            path = destination / member.name
            if not path.resolve().is_relative_to(destination.resolve()):
                raise ValueError("archive_path_escape")
            path.parent.mkdir(parents=True, exist_ok=True)
            source = archive.extractfile(member)
            if source is None:
                raise ValueError("missing_archive_member")
            path.write_bytes(source.read())


def pdf(
    pages: list[list[str]], *, clipped: bool = False, scanned: bool = False
) -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.utils import ImageReader
    import reportlab

    output = io.BytesIO()
    canvas = Canvas(output, pagesize=(612, 792), invariant=1, pageCompression=1)
    for lines in pages:
        if scanned:
            bitmap = Image.new("RGB", (1224, 1584), "white")
            draw = ImageDraw.Draw(bitmap)
            font = ImageFont.truetype(
                str(Path(reportlab.__file__).parent / "fonts/Vera.ttf"), 28
            )
            for index, line in enumerate(lines):
                draw.text((90, 100 + index * 50), line, font=font, fill="black")
            canvas.drawImage(ImageReader(bitmap), 0, 0, width=612, height=792)
        else:
            canvas.setFont("Helvetica", 11)
            for index, line in enumerate(lines):
                canvas.drawString(590 if clipped else 45, 740 - index * 20, line)
        canvas.showPage()
    canvas.save()
    return output.getvalue()


def office_bytes(path: Path) -> bytes:
    normalize_zip(path)
    return path.read_bytes()


def docx(
    path: Path,
    paragraphs: list[str],
    tables: list[list[list[str]]] | None = None,
    headings: bool = False,
    *,
    clipped: bool = False,
) -> bytes:
    from docx.shared import Inches as DocInches

    document = Document()
    if clipped:
        document.sections[0].page_width = DocInches(2)
        document.sections[0].left_margin = DocInches(1)
        document.sections[0].right_margin = DocInches(1)
    for index, text in enumerate(paragraphs):
        document.add_paragraph(
            text, style="Heading 1" if headings and index == 0 else "Normal"
        )
    for values in tables or []:
        table = document.add_table(rows=0, cols=len(values[0]))
        table.style = "Table Grid"
        for row in values:
            for cell, value in zip(table.add_row().cells, row, strict=True):
                cell.text = value
    document.save(str(path))
    return office_bytes(path)


def xlsx(
    path: Path,
    sheets: dict[str, list[list[Any]]],
    caches: dict[str, str] | None = None,
    *,
    clipped: bool = False,
) -> bytes:
    from lxml import etree

    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = (
                3 if clipped else 28
            )
            for cell in column:
                cell.alignment = Alignment(wrap_text=not clipped, vertical="top")
        for row in sheet.rows:
            sheet.row_dimensions[row[0].row].height = 32
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(path)
    if caches:
        out = io.BytesIO()
        with ZipFile(path) as source, ZipFile(out, "w", ZIP_DEFLATED) as target:
            for name in source.namelist():
                raw = source.read(name)
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                    root = etree.fromstring(raw)
                    ns = {
                        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                    }
                    for cell in root.findall(".//m:c", ns):
                        key = name.rsplit("/", 1)[1] + ":" + str(cell.get("r"))
                        if key in caches:
                            value = cell.find("m:v", ns)
                            if value is None:
                                value = etree.SubElement(cell, "{" + ns["m"] + "}v")
                            value.text = caches[key]
                    raw = etree.tostring(root)
                target.writestr(name, raw)
        path.write_bytes(out.getvalue())
    return office_bytes(path)


def pptx(
    path: Path, slides: list[tuple[str, list[str]]], *, clipped: bool = False
) -> bytes:
    deck = Presentation()
    for title, lines in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(
            Inches(9.8 if clipped else 0.5), Inches(0.4), Inches(8.5), Inches(6)
        )
        frame = box.text_frame
        frame.word_wrap = True
        for index, text in enumerate([title, *lines]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = text
            paragraph.font.size = Pt(24 if index == 0 else 18)
    deck.save(str(path))
    return office_bytes(path)


def csv_bytes(rows: list[list[Any]]) -> bytes:
    out = io.StringIO(newline="")
    csv.writer(out, lineterminator="\r\n").writerows(rows)
    return out.getvalue().encode("utf-8")


def zip_bytes(files: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as target:
        for name, raw in sorted(files.items()):
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, raw)
    return out.getvalue()
