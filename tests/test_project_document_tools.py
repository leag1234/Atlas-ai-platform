"""REQ-PROJ-05: located, inert, bounded Work document operations on real formats."""

from __future__ import annotations

import asyncio
import base64
import io
import json
import os
from pathlib import Path
import socket
import subprocess
import tempfile
import unittest
from typing import Any
from unittest.mock import AsyncMock, Mock, patch
from uuid import uuid4
import zipfile

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from packages.attachment_reader import read_structured
from services.orchestrator import code_tools, code_worker, project_document_tools
from services.orchestrator.code_workspace import Workspace, volume
from services.orchestrator.terminal_stack import IMAGE
from services.retrieval.project_identity import ProjectError


def docx_bytes(text: str = "First paragraph", table: bool = True) -> bytes:
    document = Document()
    document.add_paragraph(text)
    document.add_paragraph("Second paragraph")
    if table:
        document.add_table(rows=1, cols=2).rows[0].cells[1].text = "Table evidence"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes() -> bytes:
    output = io.BytesIO()
    pdf = canvas.Canvas(output, invariant=1)
    pdf.drawString(50, 750, "First page evidence")
    pdf.showPage()
    pdf.drawString(50, 750, "Second page evidence")
    pdf.save()
    return output.getvalue()


def xlsx_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Costs"
    sheet["A1"] = "Item"
    sheet["C2"] = 42
    sheet["D2"] = "=C2+1"
    book.create_sheet("Notes")["B3"] = "Other sheet"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def input_bytes(raw: bytes, extension: str, **fields: Any) -> dict[str, Any]:
    return dict(
        file_id=str(uuid4()),
        version=str(uuid4()),
        content=base64.b64encode(raw).decode(),
        extension=extension,
        **fields,
    )


def package_change(raw: bytes, name: str, content: bytes) -> bytes:
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as source,
        zipfile.ZipFile(output, "w") as destination,
    ):
        for entry in source.infolist():
            if entry.filename != name:
                destination.writestr(entry, source.read(entry))
        destination.writestr(name, content)
    return output.getvalue()


class ProjectDocumentToolsTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def test_pdf_docx_xlsx_precise_locators_and_no_network(self) -> None:
        with patch.object(
            socket, "create_connection", side_effect=AssertionError("network forbidden")
        ):
            pdf = read_structured(pdf_bytes(), ".pdf", dict(page_start=2, page_end=2))
            self.assertEqual(pdf["segments"][0]["locator"], {"page": 2})
            self.assertIn("Second page evidence", pdf["segments"][0]["text"])
            docx = read_structured(docx_bytes(), ".docx")
            self.assertEqual(docx["segments"][1]["locator"], {"paragraph": 2})
            self.assertEqual(
                docx["segments"][-1]["locator"],
                dict(table=1, row=1, column=2, paragraph=1),
            )
            self.assertEqual(docx["segments"][-1]["text"], "Table evidence")
            xlsx = read_structured(
                xlsx_bytes(), ".xlsx", dict(sheet="Costs", range="C2:D2")
            )
            self.assertEqual(
                xlsx["segments"][0]["locator"], dict(sheet="Costs", cell="C2")
            )
            self.assertEqual(xlsx["segments"][0]["text"], "42")
            self.assertEqual(xlsx["segments"][1]["formula"], "C2+1")
            self.assertIsNone(xlsx["segments"][1]["cached_value"])
            self.assertIn("not evaluated", str(xlsx["warnings"]))

    def test_comparison_uses_two_inputs_and_locates_changes(self) -> None:
        left = input_bytes(docx_bytes("Before"), ".docx")
        right = input_bytes(docx_bytes("After"), ".docx")
        reply = code_worker.perform(
            self.root, "compare_documents", dict(left=left, right=right)
        )
        self.assertTrue(reply["complete"])
        self.assertEqual(reply["left"]["file_id"], left["file_id"])
        self.assertEqual(reply["right"]["file_id"], right["file_id"])
        self.assertEqual(reply["changes"][0]["left"][0]["text"], "Before")
        self.assertEqual(reply["changes"][0]["right"][0]["text"], "After")
        self.assertEqual(reply["changes"][0]["left"][0]["locator"], dict(paragraph=1))
        same = code_worker.perform(
            self.root, "compare_documents", dict(left=left, right=left)
        )
        self.assertEqual(same["changes"], [])

    def test_coverage_cursor_binding_and_partial_compare(self) -> None:
        raw = "\n".join("line " + str(i) for i in range(205)).encode()
        first = read_structured(raw, ".txt")
        self.assertEqual(first["coverage"]["returned_units"], 200)
        self.assertFalse(first["coverage"]["complete"])
        second = read_structured(raw, ".txt", cursor=first["next_cursor"])
        self.assertEqual(second["coverage"]["offset"], 200)
        self.assertEqual(second["coverage"]["returned_units"], 5)
        self.assertIsNone(second["next_cursor"])
        with self.assertRaisesRegex(ValueError, "cursor"):
            read_structured(raw + b" changed", ".txt", cursor=first["next_cursor"])
        result = code_worker.perform(
            self.root,
            "compare_documents",
            dict(left=input_bytes(raw, ".txt"), right=input_bytes(raw, ".txt")),
        )
        self.assertFalse(result["complete"])
        self.assertEqual(result["changes"], [])
        with self.assertRaisesRegex(ValueError, "Located unit"):
            read_structured(b"a" * 24001, ".txt")

    def test_malformed_packages_bombs_macros_external_links_and_entities_rejected(
        self,
    ) -> None:
        raw = docx_bytes()
        malicious = [
            b"not a package",
            package_change(raw, "../escape", b"no"),
            package_change(raw, "word/vbaProject.bin", b"macro"),
            package_change(
                raw, "word/document.xml", b'<!DOCTYPE a [<!ENTITY x "boom">]><a>&x;</a>'
            ),
            package_change(raw, "word/document.xml", b"<broken>"),
            package_change(
                raw,
                "word/_rels/document.xml.rels",
                b'<Relationships><Relationship Id="r" Target="https://example.invalid/a" Type="externalData" TargetMode="External"/></Relationships>',
            ),
        ]
        for content in malicious:
            with self.subTest(size=len(content)):
                reply = code_worker.perform(
                    self.root, "read_document", input_bytes(content, ".docx")
                )
                self.assertIn("error", reply)
                self.assertNotIn("boom", reply["error"])
        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"x" * (64 * 1024 * 1024 + 1))
        result = code_worker.perform(
            self.root, "read_document", input_bytes(bomb.getvalue(), ".docx")
        )
        self.assertIn("Expanded document bytes", result["error"])
        self.assertFalse((self.root.parent / "escape").exists())

    def test_conversion_contents_validation_and_no_overwrite(self) -> None:
        original = self.root / "original.docx"
        source_bytes = docx_bytes()
        original.write_bytes(source_bytes)
        args = input_bytes(original.read_bytes(), ".docx", target="text")
        result = code_worker.perform(self.root, "convert_document", args)
        self.assertNotIn("error", result)
        text = (self.root / result["path"]).read_text()
        self.assertIn("First paragraph", text)
        self.assertIn("Table evidence", text)
        self.assertFalse(result["validation"]["published"])
        pdf = code_worker.perform(
            self.root, "convert_document", dict(args, target="pdf")
        )
        self.assertNotIn("error", pdf)
        self.assertIn(
            "Second paragraph",
            str(read_structured((self.root / pdf["path"]).read_bytes(), ".pdf")),
        )
        self.assertEqual(original.read_bytes(), source_bytes)
        refused = code_worker.perform(
            self.root,
            "convert_document",
            dict(args, target="pdf", options=dict(fidelity="lossless")),
        )
        self.assertIn("Unsupported", refused["error"])
        unicode_pdf = code_worker.perform(
            self.root,
            "convert_document",
            input_bytes(docx_bytes("日本語"), ".docx", target="pdf"),
        )
        self.assertIn("fidelity", unicode_pdf["error"])

    def test_csv_sparse_coordinates_cached_values_and_inert_cells(self) -> None:
        raw = xlsx_bytes()
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            sheet = archive.read("xl/worksheets/sheet1.xml").replace(
                b"<f>C2+1</f><v></v>", b"<f>C2+1</f><v>43</v>"
            )
        raw = package_change(raw, "xl/worksheets/sheet1.xml", sheet)
        args = input_bytes(raw, ".xlsx", target="csv", selector=dict(sheet="Costs"))
        result = code_worker.perform(self.root, "convert_document", args)
        self.assertNotIn("error", result)
        self.assertEqual((self.root / result["path"]).read_text(), "Item,,,\n,,42,43\n")
        self.assertIn(
            "exactly one",
            code_worker.perform(self.root, "convert_document", dict(args, selector={}))[
                "error"
            ],
        )
        with patch.object(code_worker, "MAX_BYTES", 1):
            self.assertIn(
                "limit 1",
                code_worker.perform(self.root, "convert_document", args)["error"],
            )

    def test_scoped_declarations_and_real_dispatch_to_worker(self) -> None:
        def names(scoped: bool) -> set[str]:
            result = set()
            for item in code_tools.declarations(scoped=scoped):
                function = item["function"]
                assert isinstance(function, dict)
                result.add(str(function["name"]))
            return result

        self.assertTrue(set(project_document_tools.SCHEMAS) <= names(True))
        self.assertFalse(set(project_document_tools.SCHEMAS) & names(False))
        raw = docx_bytes()
        import hashlib

        file_id, version = str(uuid4()), str(uuid4())
        execution = Mock()
        execution.check = AsyncMock()
        execution.session = dict(
            base_file_versions=[dict(file_id=file_id, version_id=version)]
        )
        execution.broker.public = AsyncMock(
            return_value=(
                200,
                dict(
                    version_id=version,
                    relative_path="input.docx",
                    size_bytes=len(raw),
                    sha256=hashlib.sha256(raw).hexdigest(),
                ),
            )
        )
        execution.broker.content = AsyncMock(return_value=raw)
        tools = Mock(execution=execution)
        tools.workspace.invoke = AsyncMock(
            side_effect=lambda name, args, timeout: code_worker.perform(
                self.root, name, args
            )
        )
        reply = asyncio.run(
            project_document_tools.execute(
                tools,
                "read_document",
                json.dumps(dict(file_id=file_id, version=version)),
                30,
            )
        )
        self.assertIn("First paragraph", str(reply))
        with self.assertRaisesRegex(ProjectError, "not_found"):
            asyncio.run(
                project_document_tools.execute(
                    tools,
                    "read_document",
                    json.dumps(dict(file_id=str(uuid4()), version=version)),
                    30,
                )
            )
        execution.broker.call = AsyncMock(
            return_value=dict(mode="preview", operations=[], files=[])
        )
        execution.fence.return_value = dict(
            session_id=str(uuid4()), request_id=str(uuid4()), generation=1
        )
        move = dict(
            file_id=file_id,
            base_version=version,
            expected_revision=0,
            folder_id=str(uuid4()),
            path="renamed.docx",
        )
        asyncio.run(
            project_document_tools.execute(
                tools,
                "organize_files",
                json.dumps(dict(mode="preview", operations=[move])),
                30,
            )
        )
        self.assertEqual(execution.broker.call.call_args.args[0], "organize")
        self.assertEqual(execution.broker.call.call_args.args[1]["generation"], 1)


class ProjectDocumentContainerTests(unittest.TestCase):
    def test_real_networkless_worker_structured_read_and_compare(self) -> None:
        try:
            available = (
                subprocess.run(
                    ["docker", "image", "inspect", IMAGE],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                    timeout=15,
                ).returncode
                == 0
            )
        except (OSError, subprocess.TimeoutExpired):
            available = False
        if not available:
            self.skipTest(
                "real document worker image unavailable; no image pull in task I"
            )
        with (
            tempfile.TemporaryDirectory() as temporary,
            patch.dict(
                os.environ,
                {"ATLAS_CODE_STATE_DIR": temporary, "ATLAS_PROJECT_MODE": "1"},
            ),
        ):
            workspace = Workspace(
                str(uuid4()),
                org=str(uuid4()),
                project=str(uuid4()),
                workspace=str(uuid4()),
                generation=1,
            )

            async def run() -> None:
                await workspace.prepare()
                try:
                    left = input_bytes(docx_bytes("Before"), ".docx")
                    right = input_bytes(docx_bytes("After"), ".docx")
                    read = await workspace.invoke("read_document", left, 60)
                    self.assertNotIn("error", read)
                    self.assertEqual(read["segments"][0]["text"], "Before")
                    comparison = await workspace.invoke(
                        "compare_documents", dict(left=left, right=right), 60
                    )
                    self.assertTrue(comparison["complete"])
                    self.assertEqual(
                        comparison["changes"][0]["right"][0]["text"], "After"
                    )
                    converted = await workspace.invoke(
                        "convert_document", dict(left, target="pdf"), 60
                    )
                    self.assertNotIn("error", converted)
                    self.assertTrue(converted["validation"]["readable"])
                finally:
                    await workspace.cleanup()

            try:
                asyncio.run(run())
            finally:
                subprocess.run(
                    ["docker", "volume", "rm", volume()],
                    capture_output=True,
                    check=False,
                )
