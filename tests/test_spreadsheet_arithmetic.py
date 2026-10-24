"""Owner regression, 2026-10-23: computed sales evidence reaches both answer routes."""

import base64
from decimal import Decimal
import io
import json
import os
from pathlib import Path
import subprocess
import tempfile
from typing import Any
import unittest
from unittest.mock import AsyncMock, patch
import zipfile

from openpyxl import Workbook

from services.orchestrator.documents import Attachment, extract_document
from services.orchestrator.chat_schema import ChatMessage, ChatRequest
from services.orchestrator.interactions import Interaction
from services.orchestrator.loop import Reservation, Turn
from services.orchestrator.model import Configuration, GatewayModel

SALES = Path("tests/fixtures/synthetic-sales-2026.xlsx")
QUESTION = "Summarize this sales spreadsheet: totals by region, the best month, and one chart-worthy trend."
PHRASINGS = (
    QUESTION,
    "Who sold the most?",
    "Where did we earn most?",
    "Which period stands out?",
    "SUMMARIZE REGIONAL SALES",
    "Resume les ventes par region",
)
EXPECTED = {
    "North": ("1600", "233110"),
    "South": ("1164", "205390"),
    "East": ("901", "161475"),
}


def attachment(name: str, raw: bytes) -> Attachment:
    return Attachment(filename=name, file_data=base64.b64encode(raw).decode())


class ArithmeticTests(unittest.TestCase):
    def test_exact_owner_sales_totals(self) -> None:
        result = extract_document(attachment(SALES.name, SALES.read_bytes()))
        self.assertEqual(len(result.computed_tables), 1)
        table = result.computed_tables[0]
        self.assertEqual(table["rows"], 36)
        groups = table["groups"]["Region"]
        for name, (units, revenue) in EXPECTED.items():
            self.assertEqual(groups[name]["Units"]["sum"], units)
            self.assertEqual(groups[name]["Revenue (EUR)"]["sum"], revenue)
        self.assertEqual(table["leaders_by_sum"]["Region"]["Revenue (EUR)"], ["North"])
        self.assertEqual(
            table["leaders_by_sum"]["Month"]["Revenue (EUR)"], ["February"]
        )
        self.assertEqual(
            table["groups"]["Month"]["February"]["Revenue (EUR)"]["sum"], "134465"
        )
        self.assertEqual(table["totals"]["Revenue (EUR)"]["sum"], "599975")

    def test_csv_decimal_quoted_groups_and_missing_cells(self) -> None:
        result = extract_document(
            attachment(
                "sales.csv",
                b'Region,Revenue,Units\n"North, East",0.1,2\n"North, East",0.2,\nSouth,-0.1,3\n',
            )
        )
        table = result.computed_tables[0]
        self.assertEqual(table["totals"]["Revenue"]["sum"], "0.2")
        self.assertEqual(
            table["groups"]["Region"]["North, East"]["Revenue"]["sum"], "0.3"
        )
        self.assertEqual(table["totals"]["Units"]["count"], 2)
        self.assertEqual(table["totals"]["Units"]["missing"], 1)

    def test_multiple_sheets_remain_separate(self) -> None:
        book = Workbook()
        first = book.active
        assert first is not None
        first.title = "First"
        for sheet, value in ((first, 10), (book.create_sheet("Second"), 99)):
            sheet.append(["Region", "Units"])
            sheet.append(["North", value])
        buffer = io.BytesIO()
        book.save(buffer)
        tables = extract_document(
            attachment("book.xlsx", buffer.getvalue())
        ).computed_tables
        self.assertEqual(
            [(t["sheet"], t["totals"]["Units"]["sum"]) for t in tables],
            [("First", "10"), ("Second", "99")],
        )

    def test_ods_uses_stored_numbers_and_repeated_rows(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr(
                "content.xml",
                """<office:document-content
              xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
              xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
              xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
              <table:table table:name="Sales"><table:table-row>
              <table:table-cell><text:p>Region</text:p></table:table-cell>
              <table:table-cell><text:p>Revenue</text:p></table:table-cell>
              </table:table-row><table:table-row table:number-rows-repeated="3">
              <table:table-cell><text:p>North</text:p></table:table-cell>
              <table:table-cell office:value-type="currency" office:value="12.25"><text:p>12,25 EUR</text:p></table:table-cell>
              </table:table-row></table:table></office:document-content>""",
            )
        table = extract_document(
            attachment("book.ods", buffer.getvalue())
        ).computed_tables[0]
        self.assertEqual(table["rows"], 3)
        self.assertEqual(table["groups"]["Region"]["North"]["Revenue"]["sum"], "36.75")

    def test_formula_and_mixed_columns_are_never_partial_totals(self) -> None:
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.append(["Region", "Units", "Revenue", "Mixed"])
        sheet.append(["North", 1, 10, 20])
        sheet.append(["North", 2, "=SUM(B2:B3)", "unknown"])
        buffer = io.BytesIO()
        book.save(buffer)
        table = extract_document(
            attachment("book.xlsx", buffer.getvalue())
        ).computed_tables[0]
        self.assertEqual(table["totals"]["Units"]["sum"], "3")
        for column in ("Revenue", "Mixed"):
            self.assertNotIn(column, table["totals"])
            self.assertIn(column, table["excluded_columns"])

    def test_csv_dialects_and_multiline_labels(self) -> None:
        from packages.table_arithmetic import computed_tables

        for delimiter in (",", ";", "\t"):
            raw = f'Region{delimiter}Value\n"North\nEast"{delimiter}2.5\nSouth{delimiter}-1.25\n'.encode()
            table = computed_tables(raw, ".csv")[0]
            self.assertEqual(table["totals"]["Value"]["sum"], "1.25")
            self.assertEqual(table["totals"]["Value"]["mean"], "0.625")
            self.assertEqual(table["totals"]["Value"]["min"], "-1.25")
            self.assertEqual(table["totals"]["Value"]["max"], "2.5")
            self.assertIn("North\nEast", table["groups"]["Region"])

    def test_invalid_and_oversized_tables_fail_closed(self) -> None:
        from packages.table_arithmetic import computed_tables

        for raw in (
            b"Region,Units,Units\nNorth,1,2\n",
            b"Region,Units\nNorth,1,2\n",
            b"",
            b"Region,Units\n" + b"North,1\n" * 5001,
        ):
            with self.subTest(size=len(raw)):
                result = computed_tables(raw, ".csv")[0]
                self.assertEqual(result["status"], "unavailable")
                self.assertNotIn("totals", result)
        with patch("packages.table_arithmetic.MAX_EVIDENCE_BYTES", 10):
            result = computed_tables(b"Region,Units\nNorth,1\n", ".csv")[0]
            self.assertIn("limit 10", result["reason"])
        table = computed_tables(
            b"Region,Units\n" + b"".join(f"R{i},1\n".encode() for i in range(51)),
            ".csv",
        )[0]
        self.assertEqual(table["totals"]["Units"]["sum"], "51")
        self.assertNotIn("Region", table["groups"])
        self.assertIn("limit 50", table["excluded_columns"]["Region"])

    def test_xls_refusal_provides_conversion_guidance(self) -> None:
        from services.orchestrator.documents import DocumentError
        from services.orchestrator.attachment_reading import failure_notice

        with self.assertRaisesRegex(DocumentError, "export to XLSX or CSV"):
            extract_document(attachment("legacy.xls", b"legacy binary data"))
        notice = failure_notice(
            [
                {
                    "filename": "legacy.xls",
                    "errors": [
                        "Legacy XLS computation is unavailable; export to XLSX or CSV"
                    ],
                }
            ],
            "en",
        )
        self.assertIn("export to XLSX or CSV", notice)

    def test_sandbox_archive_carries_computation_and_survives_excerpt_selection(
        self,
    ) -> None:
        from services.orchestrator.attachment_reading import (
            decode_output,
            reader_command,
        )

        with tempfile.TemporaryDirectory(dir="BRAIN") as directory:
            path = Path(directory).resolve() / "sales.csv"
            path.write_text(
                "Region,Units,Notes\n" + ("North,1," + "context " * 30 + "\n") * 100
            )
            completed = subprocess.run(
                reader_command(str(path), "total units"),
                shell=True,
                cwd=directory,
                capture_output=True,
                text=True,
                check=True,
            )
            source = decode_output(
                {
                    "exit_code": 0,
                    "output": [{"type": "output", "data": completed.stdout}],
                }
            )
            self.assertTrue(source["hierarchical_synthesis"])
            tables: Any = source["computed_tables"]
            self.assertEqual(tables[0]["totals"]["Units"]["sum"], "100")


class AnswerRouteTests(unittest.IsolatedAsyncioTestCase):
    async def test_owner_question_computed_evidence_and_answer_on_both_routes(
        self,
    ) -> None:
        from services.orchestrator.chat_pipeline import process

        for terminal, question in (
            (terminal, question) for terminal in ("0", "1") for question in PHRASINGS
        ):
            with self.subTest(terminal=terminal, question=question):
                model = GatewayModel(
                    "http://unused",
                    Configuration(
                        input_eur_per_mtok=Decimal(0),
                        output_eur_per_mtok=Decimal(0),
                        max_tokens=1000,
                    ),
                )
                request = ChatRequest(
                    messages=[ChatMessage(role="user", content=question)],
                    documents=[
                        attachment(SALES.name, SALES.read_bytes()),
                        attachment(
                            "separate.csv", b"Region,Units\nNorth,0.1\nNorth,0.2\n"
                        ),
                    ],
                )
                calls = []

                async def complete(
                    messages: list[dict[str, object]], timeout: float
                ) -> Turn:
                    calls.append(messages)
                    systems = "\n".join(
                        str(m["content"]) for m in messages if m["role"] == "system"
                    )
                    self.assertIn("computed_tables", systems)
                    self.assertIn("do not calculate totals", systems)
                    sources = next(
                        json.loads(str(m["content"]))["untrusted_attachments"]
                        for m in messages
                        if m["role"] == "user"
                        and str(m["content"]).startswith('{"untrusted_attachments"')
                    )
                    table = sources[0]["computed_tables"][0]
                    self.assertEqual(
                        sources[1]["computed_tables"][0]["totals"]["Units"]["sum"],
                        "0.3",
                    )
                    groups = table["groups"]["Region"]
                    for region, (units, revenue) in EXPECTED.items():
                        self.assertEqual(groups[region]["Units"]["sum"], units)
                        self.assertEqual(
                            groups[region]["Revenue (EUR)"]["sum"], revenue
                        )
                    # Test double renders only the evidence it actually receives.
                    answer = "; ".join(
                        f"{region}: {groups[region]['Units']['sum']} units / EUR {groups[region]['Revenue (EUR)']['sum']}"
                        for region in EXPECTED
                    )
                    leader = table["leaders_by_sum"]["Region"]["Revenue (EUR)"][0]
                    best_month = table["leaders_by_sum"]["Month"]["Revenue (EUR)"][0]
                    months = table["groups"]["Month"]
                    return Turn(
                        answer
                        + f". {leader} leads revenue. Best month: {best_month}, EUR {months[best_month]['Revenue (EUR)']['sum']}. Chart: monthly revenue, January EUR {months['January']['Revenue (EUR)']['sum']}, June EUR {months['June']['Revenue (EUR)']['sum']}.",
                        usage=Reservation(1, Decimal(0)),
                    )

                with (
                    patch.dict(
                        os.environ,
                        {"ATLAS_TERMINAL_ENABLED": terminal, "GPU_LOCAL": "0"},
                    ),
                    patch.object(
                        GatewayModel, "connect", AsyncMock(return_value=model)
                    ),
                    patch.object(model, "complete", side_effect=complete),
                ):
                    item = Interaction()
                    await process(request, item)
                self.assertTrue(calls)
                self.assertEqual(item.state, "done")
                for region, (units, revenue) in EXPECTED.items():
                    self.assertIn(
                        f"{region}: {units} units / EUR {revenue}", item.reponse
                    )
                self.assertIn("North leads revenue", item.reponse)
                self.assertIn("February, EUR 134465", item.reponse)
                self.assertIn("January EUR 86215, June EUR 109525", item.reponse)

    async def test_j44_style_workbook_request_keeps_file_tools(self) -> None:
        """Computed table evidence supports an artifact; it never replaces one."""
        from services.orchestrator.chat_pipeline import process

        model = GatewayModel(
            "http://unused",
            Configuration(
                input_eur_per_mtok=Decimal(0),
                output_eur_per_mtok=Decimal(0),
                max_tokens=1000,
            ),
        )
        request = ChatRequest(
            messages=[
                ChatMessage(
                    role="user",
                    content=(
                        "Dans cette feuille de calcul, extrais les achats du mois de "
                        "mars, compare-les avec ceux du mois de mars de l'année "
                        "précédente, et crée un nouvel onglet avec les chiffres "
                        "comparés et un graphique d'évolution"
                    ),
                )
            ],
            documents=[attachment(SALES.name, SALES.read_bytes())],
        )

        async def complete(messages: list[dict[str, object]], timeout: float) -> Turn:
            systems = "\n".join(
                str(message["content"])
                for message in messages
                if message["role"] == "system"
            )
            offered = {
                str(function["name"])
                for tool in model.available_tools(messages)
                if isinstance(function := tool.get("function"), dict)
            }
            self.assertTrue({"terminal_command", "publish_document"}.issubset(offered))
            self.assertIn("not a replacement for the requested file", systems)
            self.assertIn("use the file tools", systems)
            self.assertNotIn("do not create files", systems.lower())
            return Turn("Le classeur sera fourni.", usage=Reservation(1, Decimal(0)))

        with (
            patch.dict(os.environ, {"ATLAS_TERMINAL_ENABLED": "1", "GPU_LOCAL": "0"}),
            patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
            patch.object(model, "complete", side_effect=complete),
        ):
            item = Interaction()
            await process(request, item)
        self.assertEqual(item.state, "done")

    async def test_multisheet_and_reduced_text_keep_computed_evidence(self) -> None:
        from services.orchestrator.chat_pipeline import process

        book = Workbook()
        first = book.active
        assert first is not None
        first.title = "First"
        first.append(["Region", "Units", "Notes"])
        for _ in range(400):
            first.append(["North", 2, "long text " * 150])
        second = book.create_sheet("Second")
        second.append(["Region", "Units"])
        second.append(["North", 3])
        buffer = io.BytesIO()
        book.save(buffer)
        for terminal in ("0", "1"):
            model = GatewayModel(
                "http://unused",
                Configuration(
                    input_eur_per_mtok=Decimal(0),
                    output_eur_per_mtok=Decimal(0),
                    max_tokens=100,
                ),
            )

            async def complete(
                messages: list[dict[str, object]], timeout: float
            ) -> Turn:
                payload = next(
                    json.loads(str(m["content"]))
                    for m in messages
                    if m["role"] == "user"
                    and str(m["content"]).startswith('{"')
                    and "computed_tables" in str(m["content"])
                )
                sources = next(iter(payload.values()))
                tables = sources[0]["computed_tables"]
                self.assertEqual(tables[0]["totals"]["Units"]["sum"], "800")
                self.assertEqual(tables[1]["totals"]["Units"]["sum"], "3")
                self.assertLess(len(sources[0]["text"]), 120000)
                return Turn(
                    f"First: {tables[0]['totals']['Units']['sum']}; Second: {tables[1]['totals']['Units']['sum']}.",
                    usage=Reservation(1, Decimal(0)),
                )

            request = ChatRequest(
                messages=[ChatMessage(role="user", content="Totals per sheet")],
                documents=[
                    attachment("book.xlsx", buffer.getvalue()),
                    attachment("notes.txt", b"Keep sheets separate."),
                ],
            )
            with (
                patch.dict(
                    os.environ, {"ATLAS_TERMINAL_ENABLED": terminal, "GPU_LOCAL": "0"}
                ),
                patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
                patch.object(model, "complete", side_effect=complete),
            ):
                item = Interaction()
                await process(request, item)
            self.assertIn("First: 800; Second: 3", item.reponse)
            self.assertTrue(item.documents[0]["hierarchical_synthesis"])
