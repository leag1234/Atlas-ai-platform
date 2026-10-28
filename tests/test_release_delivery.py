"""Release 0.7: formula input and publication tool discovery regressions."""

import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from packages.doc_render import render
from services.orchestrator.project_document_tools import Publish
from services.orchestrator.project_document_input import normalize
from services.orchestrator.work_tool_selection import WorkToolSelection


class ReleaseDeliveryTests(unittest.TestCase):
    def test_structured_formula_opt_in_reaches_renderer_and_cached_values(self) -> None:
        content = dict(
            title="Costs",
            sections=[
                dict(
                    heading="Expenses",
                    tables=[
                        dict(
                            headers=["Category", "Amount"],
                            rows=[
                                ["Training", "200"],
                                ["Training", "300"],
                                ["Total", "=SUM(B4:B5)"],
                            ],
                        )
                    ],
                )
            ],
        )
        for value in (content, json.dumps(content)):
            with self.subTest(encoded=isinstance(value, str)):
                args = Publish.model_validate(
                    normalize(
                        dict(
                            path="totals.xlsx", content=value, formulas=True, sources=[]
                        )
                    )
                )
                with tempfile.TemporaryDirectory() as directory:
                    path = Path(directory) / "totals.xlsx"
                    render(args.content, path, image_root=path.parent)
                    formulas = load_workbook(path, data_only=False)
                    cached = load_workbook(path, data_only=True)
                    try:
                        self.assertEqual(formulas["1"]["B6"].value, "=SUM(B4:B5)")
                        self.assertEqual(cached["1"]["B6"].value, 500)
                    finally:
                        formulas.close()
                        cached.close()

    def test_uppercase_workbook_extension_normalizes(self) -> None:
        args = Publish.model_validate(
            normalize(
                dict(
                    path="Totals.XLSX",
                    sources=[],
                    formulas=True,
                    content="| Value |\n| --- |\n| 500 |\n| =SUM(A4) |",
                )
            )
        )
        self.assertTrue(args.content.sections[0].tables[0].formulas)

    def test_verified_code_keeps_publication_discoverable_with_optional_tools(
        self,
    ) -> None:
        selection = WorkToolSelection("code")
        selection.reopen_enabled = True
        selection.memory_enabled = True
        selection.skills_enabled = True
        selection.verified = True
        names = selection.names
        self.assertIn("forge_action", names)
        self.assertLessEqual(len(names), 8)
        self.assertIn("switch_tool_group", names)
