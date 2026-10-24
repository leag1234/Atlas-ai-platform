"""REQ-PROJ-01/03/06: synthetic located ingestion and scoped HTTP acceptance."""

from __future__ import annotations

import io
import sys
from pathlib import Path
import unittest
from typing import Any
from unittest.mock import patch
import zipfile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas

from packages.attachment_reader import MAX_BYTES
from services.retrieval.project_extraction import bounded_extract, extract

from services.retrieval.project_scoped_api import store
from services.retrieval.project_files import ProjectFiles
from services.retrieval.project_documents import ScopedProjectDocuments
from services.retrieval.search import Gateway

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_project_files import ProjectFilesTests


def fixtures() -> dict[str, bytes]:
    result: dict[str, bytes] = {}
    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(72, 720, "The project delivery window is ten days.")
    canvas.showPage()
    canvas.drawString(72, 720, "The support window is twenty days.")
    canvas.save()
    result[".pdf"] = buffer.getvalue()
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("The delivery window is ten days.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Support"
    table.cell(0, 1).text = "Twenty days"
    document.save(buffer)
    result[".docx"] = buffer.getvalue()
    buffer = io.BytesIO()
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Sales"
    sheet.append(["Region", "Revenue"])
    sheet.append(["North", 12])
    sheet.append(["South", 30])
    book.save(buffer)
    result[".xlsx"] = buffer.getvalue()
    buffer = io.BytesIO()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    assert slide.shapes.title is not None
    slide.shapes.title.text = "Delivery in ten days"
    notes = slide.notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "Speaker note: support lasts twenty days."
    deck.save(buffer)
    result[".pptx"] = buffer.getvalue()
    result[".csv"] = b'Region,Revenue\nNorth,12\n"South, Coast",30\n'
    result[".txt"] = "Delivery: ten days. Unicode: café 日本語.".encode()
    result[".md"] = b"# Delivery\nTen days."
    result[".html"] = b"<script>SECRET_SCRIPT</script><p>Delivery: ten days.</p>"
    return result


class LocatedExtractionTests(unittest.TestCase):
    def test_every_format_has_source_locators_and_useful_text(self) -> None:
        expected = {
            ".pdf": "page",
            ".docx": "paragraph",
            ".xlsx": "sheet_range",
            ".csv": "sheet_range",
            ".pptx": "slide",
            ".txt": "text",
            ".md": "text",
            ".html": "text",
        }
        for extension, raw in fixtures().items():
            with self.subTest(extension=extension):
                value = bounded_extract(raw, extension)
                self.assertTrue(value["segments"])
                self.assertEqual(
                    value["segments"][0]["locator"]["kind"], expected[extension]
                )
                self.assertGreater(value["coverage"]["selected_units"], 0)
                text = "\n".join(s["text"] for s in value["segments"])
                self.assertIn(
                    "Revenue" if extension in {".xlsx", ".csv"} else "days", text
                )
                self.assertNotIn("SECRET_SCRIPT", text)

    def test_office_tables_notes_and_exact_aggregates(self) -> None:
        values = fixtures()
        docx = extract(values[".docx"], ".docx")
        self.assertTrue(
            any(
                s["locator"]["kind"] == "table_cell" and s["text"] == "Twenty days"
                for s in docx["segments"]
            )
        )
        slide = extract(values[".pptx"], ".pptx")
        self.assertEqual(slide["coverage"]["unit"], "slides")
        self.assertEqual(slide["segments"][0]["locator"], dict(kind="slide", slide=1))
        self.assertIn("Speaker note", slide["segments"][0]["text"])
        for extension, sheet in ((".xlsx", "Sales"), (".csv", "CSV")):
            value = extract(values[extension], extension)
            aggregate = next(
                s for s in value["segments"] if s["text"].startswith("computed_tables:")
            )
            self.assertEqual(
                aggregate["locator"],
                dict(kind="sheet_range", sheet=sheet, range="A1:B3"),
            )
            self.assertIn('"sum": "42"', aggregate["text"])

    def test_encrypted_broken_oversized_and_active_content_fail_precisely(self) -> None:
        writer = PdfWriter()
        writer.append(PdfReader(io.BytesIO(fixtures()[".pdf"])))
        writer.encrypt("synthetic-test-password")
        buffer = io.BytesIO()
        writer.write(buffer)
        for raw, extension, reason in (
            (buffer.getvalue(), ".pdf", "Encrypted PDF"),
            (b"%PDF-broken", ".pdf", "PDF"),
            (b"x" * (MAX_BYTES + 1), ".txt", "bytes"),
            (b"broken", ".xlsx", "Office archive"),
        ):
            with (
                self.subTest(extension=extension, reason=reason),
                self.assertRaisesRegex(ValueError, reason),
            ):
                bounded_extract(raw, extension)
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("word/vbaProject.bin", b"never execute")
        with self.assertRaisesRegex(ValueError, "Active package"):
            bounded_extract(buffer.getvalue(), ".docx")

    def test_scanned_page_and_truncation_are_not_complete(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)
        value = extract(buffer.getvalue(), ".pdf")
        self.assertFalse(value["coverage"]["complete"])
        self.assertEqual(value["coverage"]["total_units"], 1)
        self.assertEqual(value["coverage"]["selected_units"], 0)
        self.assertIn("OCR was not performed", str(value["coverage"]["warnings"]))
        with patch("services.retrieval.project_extraction.MAX_TEXT", 10):
            value = extract(b"one two three four five", ".txt")
        self.assertFalse(value["coverage"]["complete"])
        self.assertEqual(value["coverage"]["selected_units"], 10)
        self.assertIn("limit 10", str(value["coverage"]["warnings"]))

    def test_repeated_overlapping_text_keeps_exact_character_offsets(self) -> None:
        text = "é" * 5200
        value = extract(text.encode(), ".txt")
        self.assertEqual(
            [s["locator"]["start"] for s in value["segments"]], [0, 2200, 4400]
        )
        for segment in value["segments"]:
            locator = segment["locator"]
            self.assertEqual(text[locator["start"] : locator["end"]], segment["text"])
        text = " " * 2400 + "Second window text"
        value = extract(text.encode(), ".txt")
        self.assertEqual(value["segments"][0]["locator"]["start"], 2400)

    def test_parser_receives_no_environment_credentials(self) -> None:
        import os
        import subprocess

        with (
            patch.dict(os.environ, {"INGEST_TEST_CREDENTIAL": "synthetic-only"}),
            patch(
                "services.retrieval.project_extraction.subprocess.run",
                return_value=subprocess.CompletedProcess(
                    [], 0, b'{"segments": [], "coverage": {}}'
                ),
            ) as run,
        ):
            bounded_extract(b"text", ".txt")
        self.assertEqual(
            set(run.call_args.kwargs["env"]), {"PATH", "PYTHONDONTWRITEBYTECODE"}
        )
        self.assertNotIn("INGEST_TEST_CREDENTIAL", run.call_args.kwargs["env"])

    def test_parser_timeout_is_bounded(self) -> None:
        import subprocess

        with patch(
            "services.retrieval.project_extraction.subprocess.run",
            side_effect=subprocess.TimeoutExpired("reader", 30),
        ):
            with self.assertRaisesRegex(ValueError, "elapsed seconds reached limit 30"):
                bounded_extract(b"text", ".txt")


class ScopedIngestionTests(ProjectFilesTests):
    def test_binary_upload_extraction_embeddings_and_ledger(self) -> None:
        import os

        recorded: list[list[str]] = []

        def embed(texts: list[str], kind: str) -> tuple[list[list[float]], str]:
            recorded.append(texts)
            # The actual dispatch intent must be visible from another transaction.
            import psycopg

            with psycopg.connect(self.admin_dsn) as db:
                row = db.execute(
                    "SELECT data FROM atlas_projects.ledger_events WHERE project_id=%s ORDER BY seq DESC LIMIT 1",
                    (self.project,),
                ).fetchone()
            assert row is not None
            self.assertEqual(row[0]["kind"], "model_dispatch_intent")
            self.assertEqual(row[0]["model_dispatch"]["purpose"], "embedding")
            return [
                [1.0, float(i + 1)] for i, _ in enumerate(texts)
            ], "synthetic-vector-revision"

        folder = self.folder()
        for extension, raw in fixtures().items():
            current = self.request(
                "GET", self.project_path() + "/folders", "viewer"
            ).json()["items"]
            revision = next(f["revision"] for f in current if f["id"] == folder["id"])
            with (
                patch.dict(
                    os.environ, {"ATLAS_GATEWAY_URL": "http://embedding.invalid"}
                ),
                patch.object(Gateway, "embed", side_effect=embed),
            ):
                reply = self.upload(folder, raw, "input" + extension, revision=revision)
            self.assertEqual(reply.status_code, 201, reply.text)
            file = reply.json()["file"]
            endpoint = self.project_path() + "/files/" + file["id"]
            state = self.request("GET", endpoint, "viewer")
            self.assertEqual(state.json()["index_status"], "ready", state.text)
            extraction = self.request("GET", endpoint + "/extraction", "viewer")
            self.assertEqual(extraction.status_code, 200, extraction.text)
            self.validator.validate(
                extraction.json(), self.contract["components"]["schemas"]["Extraction"]
            )
            self.assertTrue(extraction.json()["segments"])
            with store(self.app_dsn).scope(
                self.identity("editor"), self.org, self.project
            ) as scope:
                documents = ScopedProjectDocuments(ProjectFiles(scope))
                derived = documents.derived(file["version_id"])
                assert derived is not None
                self.assertEqual(len(derived["vectors"]), len(derived["chunks"]))
                self.assertEqual(
                    derived["embedding_revision"], "synthetic-vector-revision"
                )
            # Immutable upload replay does not dispatch again.
            before = len(recorded)
            self.client.send(reply.request)
            self.assertEqual(len(recorded), before)
        self.assertEqual(len(recorded), len(fixtures()))
        ledger = self.request(
            "GET", self.ledger_path() + "&limit=100", "viewer"
        ).json()["items"]
        model_events = [e for e in ledger if e["model_dispatch"]]
        self.assertEqual(len(model_events), 2 * len(fixtures()))
        for event in model_events:
            self.validator.validate(
                event, self.contract["components"]["schemas"]["LedgerEvent"]
            )
        self.assertTrue(
            any(
                r["locator"]["kind"] == "slide"
                for e in model_events
                for r in e["resource_refs"]
            )
        )

    def test_provider_failure_and_restricted_folder_do_not_claim_an_index(self) -> None:
        import os

        folder = self.folder()
        with (
            patch.dict(os.environ, {"ATLAS_GATEWAY_URL": "http://embedding.invalid"}),
            patch.object(
                Gateway, "embed", side_effect=ValueError("gateway_unavailable")
            ),
        ):
            reply = self.upload(folder, fixtures()[".pdf"], "report.pdf")
        self.assertEqual(reply.status_code, 201, reply.text)
        file = reply.json()["file"]
        path = self.project_path() + "/files/" + file["id"]
        self.assertEqual(
            self.request("GET", path, "viewer").json()["index_status"], "failed"
        )
        detail = self.request("GET", path + "/extraction", "viewer").json()
        self.assertFalse(detail["coverage"]["complete"])
        self.assertIn("Embedding gateway failed", str(detail["coverage"]["warnings"]))
        events = self.request("GET", self.ledger_path(), "viewer").json()["items"]
        self.assertEqual(events[-1]["kind"], "model_unknown")
        self.assertEqual(events[-1]["model_dispatch"]["transmission"], "possibly_sent")
        restricted = self.folder(name="Restricted", model_use="deny")
        with (
            patch("services.retrieval.project_extraction.bounded_extract") as parser,
            patch.object(Gateway, "embed") as provider,
        ):
            reply = self.upload(restricted, fixtures()[".pdf"], "restricted.pdf")
        self.assertEqual(reply.status_code, 201, reply.text)
        self.assertEqual(reply.json()["file"]["index_status"], "excluded")
        parser.assert_not_called()
        provider.assert_not_called()

    def test_parent_restrictions_prevent_binary_extraction_and_egress(self) -> None:
        parent = self.folder(name="Parent restriction", model_use="deny")
        reply = self.request(
            "POST",
            self.project_path() + "/folders",
            "editor",
            dict(
                name="Child",
                locality="server",
                model_use="allow",
                parent_id=parent["id"],
            ),
        )
        self.assertEqual(reply.status_code, 201, reply.text)
        child = reply.json()
        with (
            patch("services.retrieval.project_extraction.bounded_extract") as parser,
            patch.object(Gateway, "embed") as provider,
        ):
            uploaded = self.upload(child, fixtures()[".pdf"], "report.pdf")
        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        self.assertEqual(uploaded.json()["file"]["index_status"], "excluded")
        parser.assert_not_called()
        provider.assert_not_called()
        self.error(
            self.request(
                "GET",
                self.project_path()
                + "/files/"
                + uploaded.json()["file"]["id"]
                + "/extraction",
                "viewer",
            ),
            403,
            "policy_denied",
        )

    def test_excluded_files_and_coverage_pagination_are_scoped(self) -> None:
        folder = self.folder()
        reply = self.upload(folder, fixtures()[".pdf"], "notes.pdf")
        self.assertEqual(reply.status_code, 201, reply.text)
        file = reply.json()["file"]
        path = self.project_path() + "/files/" + file["id"] + "/extraction"
        first = self.request("GET", path + "?limit=1", "viewer")
        self.assertEqual(first.status_code, 200, first.text)
        page = first.json()
        self.assertEqual(page["segments"][0]["locator"], dict(kind="page", page=1))
        self.assertIsNotNone(page["next_cursor"])
        self.assertFalse(page["coverage"]["complete"])
        second = self.request(
            "GET", path + "?limit=1&cursor=" + page["next_cursor"], "viewer"
        )
        self.assertEqual(
            second.json()["segments"][0]["locator"], dict(kind="page", page=2)
        )
        self.error(
            self.request(
                "GET", path + "?limit=2&cursor=" + page["next_cursor"], "viewer"
            ),
            404,
            "not_found",
        )
        self.error(
            self.request(
                "GET", path.replace(self.project, self.second_project), "admin"
            ),
            404,
            "not_found",
        )
        broken = self.upload(folder, b"%PDF-1.7\nbroken", "broken.pdf", revision=1)
        self.assertEqual(broken.status_code, 201, broken.text)
        self.assertEqual(broken.json()["file"]["index_status"], "excluded")
        detail = self.request(
            "GET",
            self.project_path()
            + "/files/"
            + broken.json()["file"]["id"]
            + "/extraction",
            "viewer",
        )
        self.assertFalse(detail.json()["coverage"]["complete"])
        self.assertIn("PDF", str(detail.json()["coverage"]["warnings"]))

    def test_recorded_chat_cites_pdf_page_and_xlsx_range(self) -> None:
        import asyncio
        import json
        import os
        from decimal import Decimal
        from pathlib import Path
        from uuid import uuid4
        from unittest.mock import AsyncMock
        from packages.profiles import PROFILES
        from services.orchestrator.chat_schema import ChatRequest
        from services.orchestrator.interactions import Interaction
        from services.orchestrator.model import Configuration, GatewayModel
        from services.orchestrator.project_transcript import process_scoped_project
        from services.retrieval.project_identity import sign

        # Recorded prose is replayed unchanged; only scoped citation identifiers
        # and the gateway envelope are rebound to these synthetic source versions.
        recording = json.loads(Path("tests/cassettes/rag.json").read_text())["records"][
            1
        ]
        source_text = (
            "HELIOS Ltd, third quarter 2026: revenue 48.3 million euros; "
            "12 percent growth versus the same quarter last year. "
            "Hardware 29.1 million euros, Software 14.2 million euros, "
            "Services 5.0 million euros."
        )
        pdf_buffer = io.BytesIO()
        canvas = Canvas(pdf_buffer)
        canvas.drawString(30, 700, source_text)
        canvas.save()
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "Sales"
        sheet.append(["Report", "Revenue"])
        sheet.append([source_text, 48.3])
        xlsx_buffer = io.BytesIO()
        book.save(xlsx_buffer)
        folder = self.folder()
        versions = []
        for revision, (name, raw) in enumerate(
            (
                ("report.pdf", pdf_buffer.getvalue()),
                ("report.xlsx", xlsx_buffer.getvalue()),
            )
        ):
            reply = self.upload(folder, raw, name, revision=revision)
            self.assertEqual(reply.status_code, 201, reply.text)
            versions.append(reply.json()["file"]["version_id"])
        conversation = self.request(
            "POST",
            self.project_path() + "/conversations",
            "viewer",
            dict(title="Report evidence"),
        ).json()["id"]
        for kind in ("page", "sheet_range"):
            captured: list[dict[str, Any]] = []
            chosen: dict[str, Any] = {}

            async def broker(
                operation: str, body: dict[str, Any], key: str | None = None
            ) -> dict[str, Any]:
                reply = self.client.post(
                    "/internal/atlas/transcript/" + operation,
                    json=body,
                    headers={
                        "X-Atlas-Project-Identity": sign(
                            self.identity("viewer", self.project), self.key.encode()
                        ),
                        "X-Atlas-Execution-Secret": os.environ[
                            "ATLAS_PROJECT_EXECUTION_SECRET"
                        ],
                        "X-Atlas-Execution-Audience": "atlas-retrieval-execution",
                        "Idempotency-Key": key or str(uuid4()),
                    },
                )
                self.assertIn(reply.status_code, (200, 201), reply.text)
                result = dict(reply.json())
                if operation == "begin":
                    captured.append(result["context"])
                    candidates = result["context"]["passages"]
                    chosen.update(
                        next(
                            c
                            for c in candidates
                            if c["locator"]["kind"] == kind
                            and (kind == "page" or ":" in c["locator"]["range"])
                        )
                    )
                return result

            async def transport(
                url: str, payload: dict[str, Any], timeout: float
            ) -> dict[str, Any]:
                answer = recording["response"]["text"].replace(
                    recording["response"]["citations"][0], chosen["chunk_id"]
                )
                return dict(
                    text=json.dumps(dict(answer=answer, facts=[])),
                    calls=[],
                    usage=dict(prompt_tokens=100, completion_tokens=100),
                    cost_eur="0.0001",
                )

            model = GatewayModel(
                "http://recorded.invalid",
                Configuration(
                    input_eur_per_mtok=Decimal("0.5"),
                    output_eur_per_mtok=Decimal("1"),
                    max_tokens=2048,
                ),
            )
            request = ChatRequest.model_validate(
                dict(
                    model=PROFILES[0],
                    org_id=self.org,
                    project_id=self.project,
                    conversation_id=conversation,
                    messages=[dict(role="user", content="What was HELIOS revenue?")],
                    file_version_ids=versions,
                    stream=False,
                )
            )
            request._project_subject, request._transcript_key = "viewer", str(uuid4())
            with (
                patch.dict(
                    os.environ,
                    {
                        "ATLAS_PROJECT_EXECUTION_SECRET": "synthetic-test-broker-key-at-least-thirty-two-bytes"
                    },
                ),
                patch(
                    "services.orchestrator.project_transcript.TranscriptBroker.call",
                    side_effect=broker,
                ),
                patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
                patch.object(model, "post", side_effect=transport),
            ):
                item = Interaction()
                asyncio.run(process_scoped_project(request, item))
            self.assertEqual(item.state, "done")
            self.assertIn("48.3 million euros", item.reponse)
            self.assertIn("/sources/" + chosen["chunk_id"], item.reponse)
            resolved = self.request(
                "GET", self.project_path() + "/sources/" + chosen["chunk_id"], "viewer"
            )
            self.assertEqual(resolved.status_code, 200, resolved.text)
            self.assertEqual(resolved.json()["locator"]["kind"], kind)
            self.assertIn("48.3", resolved.json()["text"])
            self.assertEqual(
                {c["version_id"] for c in captured[0]["passages"]}, set(versions)
            )


def load_tests(
    loader: unittest.TestLoader, tests: unittest.TestSuite, pattern: str | None
) -> unittest.TestSuite:
    # Reuse database fixtures, not a second copy of their inherited test cases.
    return unittest.TestSuite(
        [
            loader.loadTestsFromTestCase(LocatedExtractionTests),
            unittest.TestSuite(
                ScopedIngestionTests(name)
                for name in ScopedIngestionTests.__dict__
                if name.startswith("test_")
            ),
        ]
    )
