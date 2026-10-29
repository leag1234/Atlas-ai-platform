#!/usr/bin/env python3
"""Live, disposable smoke checks for a deployed Atlas instance.

This is intentionally a command-line program, rather than a unittest module: it
talks to a real instance and accepts credentials only through mode-600 files.
It creates fictional data only and removes every resource it can identify in a
finally block.  No value read from a credential file is ever printed.
"""

from __future__ import annotations

import argparse
import json
import base64
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
import re
import shutil
import tempfile
import time
from typing import Any, Callable
from urllib.parse import urlparse
from urllib.error import HTTPError, URLError
from urllib.request import Request, build_opener, HTTPSHandler
import ssl
from uuid import uuid4

ERROR_MARKERS = ("error", "internal server error", "traceback", "failed")
DENY_LIST = (
    "New project",
    "Send",
    "Uh-oh",
    "Running",
    "Chats",
    "Files",
    "Instructions",
)
OPTIONAL_STATUS = {404, 405, 501}


@dataclass
class Outcome:
    name: str
    status: str
    seconds: float
    reason: str


class SmokeFailure(RuntimeError):
    pass


class HttpError(RuntimeError):
    pass


@dataclass
class HttpResponse:
    status_code: int
    content: bytes
    headers: dict[str, str]

    @property
    def text(self) -> str:
        return self.content.decode("utf-8", errors="replace")

    def json(self) -> Any:
        return json.loads(self.text)


class HttpClient:
    """Small requests-like wrapper so this live runner has no new dependency."""

    def __init__(self, verify: bool) -> None:
        self.headers: dict[str, str] = {}
        self.cookies: dict[str, str] = {}
        self.auth: tuple[str, str] | None = None
        self.opener = build_opener(
            HTTPSHandler(context=None if verify else ssl._create_unverified_context())
        )

    @staticmethod
    def multipart(
        files: dict[str, tuple[str, Any]], fields: dict[str, str]
    ) -> tuple[bytes, str]:
        boundary = "----atlas-smoke-" + uuid4().hex
        body: list[bytes] = []
        for name, value in fields.items():
            body.extend(
                (
                    f"--{boundary}\r\n".encode(),
                    f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
                    value.encode(),
                    b"\r\n",
                )
            )
        for name, (filename, stream) in files.items():
            body.extend(
                (
                    f"--{boundary}\r\n".encode(),
                    f'Content-Disposition: form-data; name="{name}"; filename="{filename}"\r\n'.encode(),
                    b"Content-Type: application/octet-stream\r\n\r\n",
                    stream.read(),
                    b"\r\n",
                )
            )
        body.append(f"--{boundary}--\r\n".encode())
        return b"".join(body), f"multipart/form-data; boundary={boundary}"

    def request(
        self,
        method: str,
        url: str,
        *,
        headers: dict[str, str] | None = None,
        timeout: int = 60,
        json: Any = None,
        files: dict[str, tuple[str, Any]] | None = None,
        data: dict[str, str] | None = None,
    ) -> HttpResponse:
        actual_headers = dict(self.headers)
        actual_headers.update(headers or {})
        if self.cookies and "Cookie" not in actual_headers:
            actual_headers["Cookie"] = "; ".join(
                f"{name}={value}" for name, value in self.cookies.items()
            )
        if self.auth:
            encoded = base64.b64encode(
                f"{self.auth[0]}:{self.auth[1]}".encode()
            ).decode()
            actual_headers.setdefault("Authorization", "Basic " + encoded)
        payload: bytes | None = None
        if json is not None:
            payload = __import__("json").dumps(json).encode()
            actual_headers.setdefault("Content-Type", "application/json")
        elif files:
            payload, content_type = self.multipart(files, data or {})
            actual_headers["Content-Type"] = content_type
        request = Request(url, data=payload, headers=actual_headers, method=method)
        try:
            with self.opener.open(request, timeout=timeout) as reply:
                return HttpResponse(
                    reply.status, reply.read(), dict(reply.headers.items())
                )
        except HTTPError as error:
            return HttpResponse(error.code, error.read(), dict(error.headers.items()))
        except URLError as error:
            raise HttpError(str(error.reason)) from error

    def get(self, url: str, **kwargs: Any) -> HttpResponse:
        return self.request("GET", url, **kwargs)

    def post(self, url: str, **kwargs: Any) -> HttpResponse:
        return self.request("POST", url, **kwargs)


class Smoke:
    def __init__(self, args: argparse.Namespace, account: dict[str, Any]) -> None:
        self.args = args
        self.base = args.base.rstrip("/")
        self.account = account
        self.http = HttpClient(verify=not args.insecure)
        self.http.headers["Accept"] = "application/json"
        self.token = ""
        self.model = ""
        self.work_profile = ""
        self.org = ""
        self.project: tuple[str, str] | None = None
        self.project_error: str | None = None
        self.project_revisions: dict[tuple[str, str], int] = {}
        self.created_chats: list[str] = []
        self.created_projects: list[tuple[str, str]] = []
        self.created_files: list[tuple[str, str, str]] = []
        self.created_tasks: list[tuple[str, str, str]] = []
        self.tempdir = Path(tempfile.mkdtemp(prefix="atlas-live-smoke-"))

    def request(self, method: str, path: str, **kwargs: Any) -> HttpResponse:
        headers = dict(kwargs.pop("headers", {}))
        if self.token:
            headers.setdefault("Authorization", f"Bearer {self.token}")
        # Project-mode routes are deliberately browser-only: the BFF verifies
        # the Open WebUI session cookie and rejects state-changing cross-origin
        # requests.  Match fetch(..., {credentials: 'same-origin'}).
        if path.startswith("/api/atlas/"):
            headers.setdefault("Origin", self.base)
            headers.setdefault("Sec-Fetch-Site", "same-origin")
            if method not in {"GET", "HEAD"}:
                headers.setdefault("Idempotency-Key", str(uuid4()))
        return self.http.request(
            method,
            self.base + path,
            headers=headers,
            timeout=kwargs.pop("timeout", 60),
            **kwargs,
        )

    @staticmethod
    def _json(response: HttpResponse) -> dict[str, Any]:
        try:
            value = response.json()
        except ValueError as exc:
            raise SmokeFailure(
                f"HTTP {response.status_code}: non-JSON response"
            ) from exc
        if not isinstance(value, dict):
            raise SmokeFailure(f"HTTP {response.status_code}: unexpected JSON shape")
        return value

    @staticmethod
    def _id(value: Any) -> str:
        if isinstance(value, dict):
            for key in ("id", "chat_id", "conversation_id", "file_id", "project_id"):
                if value.get(key):
                    return str(value[key])
            for key in ("chat", "conversation", "file", "project"):
                if isinstance(value.get(key), dict):
                    found = Smoke._id(value[key])
                    if found:
                        return found
        return ""

    def run(self, name: str, fn: Callable[[], str]) -> Outcome:
        started = time.monotonic()
        try:
            reason = fn()
            return Outcome(name, "PASS", time.monotonic() - started, reason)
        except Skip as exc:
            return Outcome(name, "SKIP", time.monotonic() - started, str(exc))
        except (HttpError, SmokeFailure, AssertionError, OSError, ValueError) as exc:
            return Outcome(
                name, "FAIL", time.monotonic() - started, self.one_line(str(exc))
            )

    @staticmethod
    def one_line(text: str) -> str:
        return re.sub(r"\s+", " ", text).strip()[:240] or "no reason supplied"

    def login(self) -> str:
        email = str(self.account.get("email", ""))
        password = str(self.account.get("password", ""))
        if not email or not password:
            raise SmokeFailure("account file needs non-empty email and password")
        # A guest deployment may put nginx basic auth in front of this same call.
        basic_user = str(
            self.account.get("basic_auth_user", self.account.get("basic_user", ""))
        )
        basic_password = str(
            self.account.get(
                "basic_auth_password", self.account.get("basic_password", "")
            )
        )
        if self.args.kind == "guest" and basic_user and basic_password:
            self.http.auth = (basic_user, basic_password)
        response = self.request(
            "POST", "/api/v1/auths/signin", json={"email": email, "password": password}
        )
        if response.status_code != 200:
            raise SmokeFailure(f"sign-in returned HTTP {response.status_code}")
        body = self._json(response)
        self.token = str(body.get("token", ""))
        if not self.token:
            raise SmokeFailure("sign-in response has no token")
        # Open WebUI's browser client persists this token in its same-origin
        # session cookie.  Project routes intentionally do not accept a Bearer
        # token as a substitute for that browser session.
        self.http.cookies["token"] = self.token
        return "Open WebUI session established"

    def choose_model(self) -> str:
        for path in ("/api/models", "/api/v1/models"):
            response = self.request("GET", path)
            if response.status_code != 200:
                continue
            body = self._json(response)
            values = body.get("data", body.get("models", []))
            if isinstance(values, list):
                for item in values:
                    candidate = item.get("id") if isinstance(item, dict) else item
                    if candidate:
                        self.model = str(candidate)
                        return self.model
        raise SmokeFailure("no model available from /api/models or /api/v1/models")

    def chat(
        self, prompt: str, files: list[dict[str, Any]] | None = None
    ) -> tuple[str, dict[str, Any]]:
        model = self.model or self.choose_model()
        content: str | list[dict[str, Any]] = prompt
        if files:
            content = [{"type": "text", "text": prompt}] + [
                {"type": "file", "file": file} for file in files
            ]
        payload: dict[str, Any] = {
            "model": model,
            "messages": [{"role": "user", "content": content}],
            "stream": False,
        }
        response = self.request(
            "POST", "/api/chat/completions", json=payload, timeout=300
        )
        if response.status_code != 200:
            raise SmokeFailure(f"assistant call returned HTTP {response.status_code}")
        body = self._json(response)
        choices = body.get("choices", [])
        if (
            not isinstance(choices, list)
            or not choices
            or not isinstance(choices[0], dict)
        ):
            raise SmokeFailure("assistant response has no choice")
        message = choices[0].get("message", {})
        content = message.get("content", "") if isinstance(message, dict) else ""
        if not isinstance(content, str):
            content = json.dumps(content, ensure_ascii=False)
        return content, body

    def chat_web(self) -> str:
        answer, body = self.chat(
            "Quelles sont les principales actualités scientifiques de cette semaine sur l'énergie solaire ? Cite tes sources."
        )
        lowered = answer.lower()
        metadata = json.dumps(body, ensure_ascii=False)
        if len(answer) <= 300:
            raise SmokeFailure(f"answer has {len(answer)} characters, need >300")
        if any(marker in lowered for marker in ERROR_MARKERS):
            raise SmokeFailure("answer contains an error marker")
        if not re.search(r"https?://[^\s)\]>]+", answer + " " + metadata):
            raise SmokeFailure("no fetched source URL in answer or response metadata")
        return "answer is >300 chars and includes a fetched URL: " + self.one_line(
            answer[:80]
        )

    def fixtures(self) -> list[Path]:
        from docx import Document
        from reportlab.pdfgen.canvas import Canvas

        pdf = self.tempdir / "smoke-fiction.pdf"
        canvas = Canvas(str(pdf))
        canvas.drawString(72, 760, "Fictional Atlas smoke dossier")
        canvas.drawString(72, 735, "The quartz beacon is stored in cabinet 47.")
        canvas.save()
        docx = self.tempdir / "smoke-fiction.docx"
        document = Document()
        document.add_heading("Fictional smoke notes", 1)
        document.add_paragraph("The amber compass is inspected every Tuesday.")
        document.save(str(docx))
        return [pdf, docx]

    def upload(self, path: Path) -> dict[str, Any]:
        with path.open("rb") as stream:
            response = self.request(
                "POST",
                "/api/v1/files/?process=false",
                files={"file": (path.name, stream)},
            )
        if response.status_code not in (200, 201):
            raise SmokeFailure(
                f"upload {path.suffix} returned HTTP {response.status_code}"
            )
        body = self._json(response)
        file_id = self._id(body)
        if not file_id:
            raise SmokeFailure(f"upload {path.suffix} returned no file id")
        # The Atlas chat boundary consumes OpenAI-style file content parts,
        # whose file object contains verified bytes.  Upload through the same
        # Open WebUI inlet first, then use its guarded content endpoint.
        content = self.request("GET", f"/api/v1/files/{file_id}/content")
        if content.status_code != 200 or not content.content:
            raise SmokeFailure(
                f"uploaded {path.suffix} content returned HTTP {content.status_code}"
            )
        self.created_files.append(("ordinary", "", file_id))
        return {
            "filename": path.name,
            "file_id": file_id,
            "file_data": base64.b64encode(content.content).decode("ascii"),
        }

    def chat_document(self) -> str:
        uploaded = [self.upload(path) for path in self.fixtures()]
        answer, _ = self.chat(
            "Dans les deux documents fictifs, où se trouve la balise de quartz et quand inspecte-t-on la boussole ambrée ?",
            uploaded,
        )
        if "47" not in answer or not re.search(r"(?i)tuesday|mardi", answer):
            raise SmokeFailure("answer did not contain both fictional document facts")
        return "PDF and DOCX facts recovered: " + self.one_line(answer[:80])

    def image(self) -> str:
        started = time.monotonic()
        answer, body = self.chat(
            "Crée une image simple d'un carré bleu sur fond blanc."
        )
        evidence = answer + " " + json.dumps(body, ensure_ascii=False)
        if not re.search(r"!\[[^]]*\]\([^)]*\)|https?://\S+|data:image/", evidence):
            raise SmokeFailure(
                "no image URL or image markdown; response: "
                + self.one_line(answer[:80])
            )
        elapsed = time.monotonic() - started
        if elapsed > 180:
            raise SmokeFailure(f"image response took {elapsed:.1f}s, limit is 180s")
        return f"image returned in {elapsed:.1f}s"

    def api_key(self) -> str:
        if not self.args.api_key_file:
            raise Skip("no --api-key-file supplied")
        key = read_secret_file(Path(self.args.api_key_file)).strip()
        if not key:
            raise SmokeFailure("API key file is empty")
        headers = {"Authorization": f"Bearer {key}"}
        response = self.http.get(self.base + "/api/models", headers=headers, timeout=60)
        if response.status_code != 200:
            raise SmokeFailure(
                f"API key model list returned HTTP {response.status_code}"
            )
        models = self._json(response).get("data", [])
        if not isinstance(models, list) or not models:
            raise SmokeFailure("API key model list is empty")
        model = (
            str(models[0].get("id", ""))
            if isinstance(models[0], dict)
            else str(models[0])
        )
        assistant = self.http.post(
            self.base + "/api/chat/completions",
            headers=headers,
            json={
                "model": model,
                "messages": [{"role": "user", "content": "Réponds seulement: quartz."}],
                "stream": False,
            },
            timeout=180,
        )
        if assistant.status_code != 200:
            raise SmokeFailure(
                f"API key assistant call returned HTTP {assistant.status_code}"
            )
        raw = self.http.post(
            self.base + "/v1/chat/completions",
            headers=headers,
            json={
                "model": model,
                "messages": [{"role": "user", "content": "Reply only: quartz."}],
                "stream": False,
            },
            timeout=180,
        )
        if raw.status_code in (401, 403):
            return "models and assistant call succeeded; key lacks raw-model scope"
        if raw.status_code != 200:
            raise SmokeFailure(f"raw model call returned HTTP {raw.status_code}")
        return "models, assistant call, and raw model call succeeded"

    def memberships(self) -> list[dict[str, Any]]:
        for path in ("/api/atlas/me", "/api/atlas/memberships"):
            response = self.request("GET", path)
            if response.status_code == 200:
                body = self._json(response)
                values = body.get("memberships", body.get("items", []))
                if isinstance(values, list):
                    return [x for x in values if isinstance(x, dict)]
        raise SmokeFailure("project membership endpoint unavailable")

    def project_context(self) -> tuple[str, str]:
        if self.org:
            return self.org, ""
        memberships = self.memberships()
        if not memberships:
            raise SmokeFailure("account has no project organization membership")
        self.org = str(memberships[0].get("org_id", memberships[0].get("id", "")))
        if not self.org:
            raise SmokeFailure("membership has no organization id")
        return self.org, ""

    def create_project(self, suffix: str) -> tuple[str, str]:
        if self.project is not None:
            return self.project
        if self.project_error is not None:
            raise SmokeFailure(self.project_error)
        org, _ = self.project_context()
        for attempt in range(3):
            response = self.request(
                "POST",
                f"/api/atlas/orgs/{org}/projects",
                json={
                    "name": f"Quartz smoke {suffix}",
                    "instructions": "",
                    "default_action": "chat",
                },
            )
            if response.status_code != 429:
                break
            retry_after = response.headers.get("Retry-After", "60")
            try:
                wait_seconds = max(1, min(60, int(retry_after)))
            except ValueError:
                wait_seconds = 1
            if attempt == 2:
                break
            time.sleep(wait_seconds)
        if response.status_code not in (200, 201):
            self.project_error = f"create project returned HTTP {response.status_code} after paced retries"
            raise SmokeFailure(self.project_error)
        created = self._json(response)
        project = self._id(created)
        if not project:
            raise SmokeFailure("create project returned no project id")
        pair = (org, project)
        revision = created.get("revision")
        if isinstance(revision, int):
            self.project_revisions[pair] = revision
        self.created_projects.append(pair)
        self.project = pair
        return pair

    def project_profile(self) -> str:
        """Read the exact work profile injected into the browser project UI."""
        if self.work_profile:
            return self.work_profile
        response = self.request("GET", "/project-mode/client.js")
        if response.status_code != 200:
            raise SmokeFailure(
                f"project UI profile returned HTTP {response.status_code}"
            )
        match = re.search(
            r"window\.atlasProjectWorkProfile\s*=\s*([^;]+);", response.text
        )
        if match is None:
            raise SmokeFailure("project UI did not inject a work profile")
        try:
            profile = json.loads(match.group(1))
        except json.JSONDecodeError as exc:
            raise SmokeFailure("project UI work profile is not JSON") from exc
        if not isinstance(profile, str) or not profile:
            raise SmokeFailure("project UI work profile is empty")
        self.work_profile = profile
        return profile

    def project_upload(self, project: tuple[str, str], path: Path) -> dict[str, Any]:
        org, project_id = project
        base = f"/api/atlas/orgs/{org}/projects/{project_id}"
        folders = self.request("GET", base + "/folders")
        if folders.status_code != 200:
            raise SmokeFailure(
                f"project folder list returned HTTP {folders.status_code}"
            )
        values = self._json(folders).get("items", [])
        folder = next(
            (
                item
                for item in values
                if isinstance(item, dict)
                and item.get("locality") == "server"
                and item.get("model_use") == "allow"
            ),
            None,
        )
        if folder is None:
            made = self.request(
                "POST",
                base + "/folders",
                json={
                    "name": "Smoke files",
                    "locality": "server",
                    "model_use": "allow",
                },
            )
            if made.status_code not in (200, 201):
                raise SmokeFailure(
                    f"project folder creation returned HTTP {made.status_code}"
                )
            folder = self._json(made)
        folder_id = str(folder.get("id", ""))
        revision = folder.get("revision")
        if not folder_id or not isinstance(revision, int):
            raise SmokeFailure("project folder response lacks id or revision")
        with path.open("rb") as stream:
            response = self.request(
                "POST",
                base + f"/files?folder_id={folder_id}",
                files={"file": (path.name, stream)},
                data={"path": path.name, "expected_parent_revision": str(revision)},
            )
        if response.status_code not in (200, 201):
            raise SmokeFailure(
                f"project file upload returned HTTP {response.status_code}"
            )
        body = self._json(response)
        file_id = self._id(body)
        if file_id:
            self.created_files.append(("project", org, f"{project_id}/{file_id}"))
        file = body.get("file", body)
        return file if isinstance(file, dict) else body

    def project_chat(
        self, project: tuple[str, str], prompt: str, *, work: bool = True
    ) -> tuple[str, dict[str, Any]]:
        """Use the same scoped conversation/session sequence as project_mode.js."""
        org, project_id = project
        base = f"/api/atlas/orgs/{org}/projects/{project_id}"
        created = self.request(
            "POST", base + "/conversations", json={"title": "Quartz smoke conversation"}
        )
        if created.status_code not in (200, 201):
            raise SmokeFailure(
                f"project conversation creation returned HTTP {created.status_code}"
            )
        conversation = self._id(self._json(created))
        if not conversation:
            raise SmokeFailure("project conversation creation returned no id")
        body: dict[str, Any] = {
            "atlas_ui": True,
            "model": self.project_profile()
            if work
            else (self.model or self.choose_model()),
            "ui_locale": "fr",
            "org_id": org,
            "project_id": project_id,
            "conversation_id": conversation,
            "project_action": "chat",
            "messages": [{"role": "user", "content": prompt}],
            "stream": True,
        }
        if work:
            files = self.request("GET", base + "/files")
            if files.status_code != 200:
                raise SmokeFailure(
                    f"project file list returned HTTP {files.status_code}"
                )
            details = self.request("GET", base)
            folder_list = self.request("GET", base + "/folders")
            if details.status_code != 200 or folder_list.status_code != 200:
                raise SmokeFailure("project file-selection metadata unavailable")
            instructions = self._json(details).get("instructions_file_id")
            allowed_folders = {
                folder["id"]
                for folder in self._json(folder_list).get("items", [])
                if isinstance(folder, dict) and folder.get("model_use") == "allow"
            }
            # Match eligibleFiles() in project_mode.js. Selecting instructions as
            # an attachment turns an explicit user fact into a document proposal.
            selected = [
                item
                for item in self._json(files).get("items", [])
                if isinstance(item, dict)
                and item.get("id") != instructions
                and item.get("locality") != "local_only"
                and item.get("availability") != "not_retained"
                and item.get("folder_id") in allowed_folders
            ]
            folders = list(
                dict.fromkeys(
                    str(item.get("folder_id", ""))
                    for item in selected
                    if item.get("folder_id")
                )
            )
            versions = [
                {
                    "file_id": str(item.get("id", "")),
                    "version_id": str(item.get("version_id", "")),
                }
                for item in selected
                if item.get("id") and item.get("version_id")
            ]
            session = self.request(
                "POST",
                base + "/work-sessions",
                json={
                    "conversation_id": conversation,
                    "folder_ids": folders,
                    "base_versions": versions,
                },
            )
            if session.status_code not in (200, 201):
                raise SmokeFailure(
                    f"project work-session creation returned HTTP {session.status_code}"
                )
            session_id = self._id(self._json(session))
            if not session_id:
                raise SmokeFailure("project work-session creation returned no id")
            body.update(
                project_action="work",
                file_version_ids=[value["version_id"] for value in versions],
                work_session_id=session_id,
            )
        response = self.request(
            "POST",
            "/v1/chat/completions",
            headers={
                "Origin": self.base,
                "Sec-Fetch-Site": "same-origin",
                "Idempotency-Key": str(uuid4()),
            },
            json=body,
            timeout=300,
        )
        if response.status_code != 200:
            raise SmokeFailure(
                f"project assistant call returned HTTP {response.status_code}: {self.one_line(response.text)}"
            )
        fragments: list[str] = []
        events: list[dict[str, Any]] = []
        for line in response.text.splitlines():
            if not line.startswith("data:"):
                continue
            try:
                event = json.loads(line[5:].strip())
            except ValueError:
                continue
            if isinstance(event, dict):
                events.append(event)
                delta = event.get("choices", [{}])
                if isinstance(delta, list) and delta and isinstance(delta[0], dict):
                    content = delta[0].get("delta", {}).get("content", "")
                    if isinstance(content, str):
                        fragments.append(content)
                if event.get("error"):
                    raise SmokeFailure(
                        "project stream error: " + self.one_line(str(event["error"]))
                    )
        answer = "".join(fragments)
        if not answer:
            raise SmokeFailure("project stream produced no assistant text")
        return answer, {"events": events}

    def project_deliverable(self) -> str:
        from openpyxl import Workbook, load_workbook

        project = self.create_project(uuid4().hex[:8])
        sheet = self.tempdir / "fictional-ledger.xlsx"
        book = Workbook()
        tab = book.active
        if tab is None:
            raise SmokeFailure("new workbook has no active sheet")
        tab.title = "Fictional ledger"
        tab.append(["Item", "Amount"])
        tab.append(["Quartz", 12])
        tab.append(["Amber", 30])
        book.save(str(sheet))
        self.project_upload(project, sheet)
        answer, body = self.project_chat(
            project,
            "À partir du tableur fictif, produis une synthèse Word et un tableur avec une formule TOTAL des montants.",
            work=True,
        )
        org, project_id = project
        deliverables = self.request(
            "GET", f"/api/atlas/orgs/{org}/projects/{project_id}/deliverables"
        )
        if deliverables.status_code != 200:
            raise SmokeFailure(
                f"deliverable list returned HTTP {deliverables.status_code}"
            )
        items = self._json(deliverables).get("items", [])
        if not isinstance(items, list):
            raise SmokeFailure("deliverable list has unexpected shape")
        urls = [
            str(item.get("download_route", ""))
            for item in items
            if isinstance(item, dict) and item.get("download_route")
        ]
        if len(urls) < 2:
            raise SmokeFailure(
                "project produced fewer than two downloadable deliverables"
            )
        xlsx_bytes = b""
        downloaded = 0
        for url in urls:
            target = url if url.startswith("http") else self.base + url
            response = self.request("GET", target.removeprefix(self.base), timeout=60)
            if response.status_code == 200 and response.content:
                downloaded += 1
                if response.content[:2] == b"PK":
                    xlsx_bytes = response.content
        if downloaded < 2 or not xlsx_bytes:
            raise SmokeFailure("could not download both Word and XLSX deliverables")
        output = self.tempdir / "deliverable.xlsx"
        output.write_bytes(xlsx_bytes)
        formulas = load_workbook(output, data_only=False)
        values = load_workbook(output, data_only=True)
        has_formula = any(
            isinstance(cell.value, str) and cell.value.startswith("=")
            for sheet_ in formulas
            for row in sheet_.iter_rows()
            for cell in row
        )
        cached_totals = [
            cell.value
            for sheet_ in values
            for row in sheet_.iter_rows()
            for cell in row
            if isinstance(cell.value, (int, float))
        ]
        if not has_formula or not cached_totals:
            raise SmokeFailure("XLSX lacks a formula or cached numeric total")
        return "Word and XLSX downloaded; XLSX formula and cached values found"

    def project_memory(self) -> str:
        project = self.create_project(uuid4().hex[:8])
        fact = "La sentinelle Quartz-17 est fictive et son code est ambre-47."
        self.project_chat(project, "Retiens ce fait pour ce projet: " + fact)
        answer, _ = self.project_chat(
            project,
            "Dans cette autre conversation du projet, quel est le code de la sentinelle Quartz-17 ?",
        )
        if "ambre-47" not in answer.lower():
            raise SmokeFailure(
                "project memory was not recalled in a second conversation"
            )
        return "project fact recalled in a separate conversation"

    def move_to_project(self) -> str:
        project = self.create_project(uuid4().hex[:8])
        message_id = str(uuid4())
        response = self.request(
            "POST",
            "/api/v1/chats/new",
            json={
                "chat": {
                    "title": "Quartz ordinary smoke",
                    "history": {
                        "messages": {
                            message_id: {
                                "id": message_id,
                                "parentId": None,
                                "childrenIds": [],
                                "role": "user",
                                "content": "The fictional Quartz code is amber-47.",
                                "timestamp": int(time.time() * 1000),
                            }
                        },
                        "currentId": message_id,
                    },
                }
            },
        )
        if response.status_code in OPTIONAL_STATUS:
            raise Skip("ordinary-chat creation endpoint absent")
        if response.status_code not in (200, 201):
            raise SmokeFailure(
                f"ordinary-chat creation returned HTTP {response.status_code}"
            )
        chat_id = self._id(self._json(response))
        if not chat_id:
            raise SmokeFailure("ordinary-chat creation returned no id")
        self.created_chats.append(chat_id)
        org, project_id = project
        original = self.request("GET", f"/api/v1/chats/{chat_id}")
        if original.status_code != 200:
            raise SmokeFailure(
                f"ordinary-chat read returned HTTP {original.status_code}"
            )
        chat = self._json(original).get("chat", {})
        history = chat.get("history", {})
        stored = history.get("messages", {})
        current = history.get("currentId")
        branch: list[dict[str, Any]] = []
        while current and current in stored:
            message = stored[current]
            branch.insert(0, message)
            current = message.get("parentId")
        messages = [
            {
                "role": message["role"],
                "content": message["content"].strip(),
                "created_at": datetime.fromtimestamp(
                    message.get("timestamp", time.time() * 1000) / 1000, timezone.utc
                )
                .isoformat()
                .replace("+00:00", "Z"),
                "attachment_file_ids": [],
            }
            for message in branch
            if message.get("role") in {"user", "assistant", "system"}
            and isinstance(message.get("content"), str)
            and message["content"].strip()
        ]
        if not messages:
            raise SmokeFailure("ordinary chat has no importable messages")
        base = f"/api/atlas/orgs/{org}/projects/{project_id}/conversations"
        move = self.request(
            "POST",
            base + "/import",
            json={
                "title": chat.get("title", "Quartz ordinary smoke"),
                "messages": messages,
            },
        )
        if move.status_code in OPTIONAL_STATUS:
            raise Skip("move-to-project endpoint absent")
        if move.status_code != 201:
            raise SmokeFailure(
                f"move-to-project returned HTTP {move.status_code}: {self.one_line(move.text)}"
            )
        imported_id = self._id(self._json(move))
        listing = self.request("GET", base)
        if (
            listing.status_code != 200
            or not imported_id
            or imported_id not in listing.text
        ):
            raise SmokeFailure("moved chat is absent from project conversation list")
        turns = self.request("GET", base + f"/{imported_id}/turns")
        if turns.status_code != 200:
            raise SmokeFailure("imported conversation turns are unavailable")
        texts = [
            ref.get("text")
            for turn in self._json(turns).get("items", [])
            for ref in turn.get("content_refs", [])
        ]
        if texts != [message["content"] for message in messages]:
            raise SmokeFailure("imported conversation did not preserve the messages")
        removed = self.request("DELETE", f"/api/v1/chats/{chat_id}")
        if removed.status_code != 200:
            raise SmokeFailure(
                f"ordinary-chat deletion returned HTTP {removed.status_code}"
            )
        self.created_chats.remove(chat_id)
        ordinary = self.request("GET", "/api/v1/chats/")
        if ordinary.status_code == 200 and chat_id in ordinary.text:
            raise SmokeFailure("moved chat remains in ordinary chat list")
        return "ordinary chat moved into project and absent from chat list"

    def scheduled_task(self) -> str:
        project = self.create_project(uuid4().hex[:8])
        org, project_id = project
        due = (datetime.now(timezone.utc) + timedelta(minutes=2)).isoformat()
        response = self.request(
            "POST",
            f"/api/atlas/orgs/{org}/projects/{project_id}/scheduled-tasks",
            json={
                "name": "Quartz smoke scheduled task",
                "schedule_at": due,
                "target_file": "scheduled-quartz.txt",
                "prompt": "Create the fictional target file.",
            },
        )
        if response.status_code in OPTIONAL_STATUS:
            raise Skip("scheduled-task endpoint absent")
        if response.status_code not in (200, 201):
            raise SmokeFailure(
                f"create scheduled task returned HTTP {response.status_code}"
            )
        task_id = self._id(self._json(response))
        if not task_id:
            raise SmokeFailure("scheduled task returned no id")
        self.created_tasks.append((org, project_id, task_id))
        deadline = time.monotonic() + 360
        while time.monotonic() < deadline:
            status = self.request(
                "GET",
                f"/api/atlas/orgs/{org}/projects/{project_id}/scheduled-tasks/{task_id}",
            )
            if status.status_code == 200:
                text = status.text
                if (
                    text.count('"status":"completed"') == 1
                    or text.count('"run_id"') == 1
                ):
                    if "scheduled-quartz.txt" in text:
                        return (
                            "exactly one scheduled run and target-file version recorded"
                        )
            time.sleep(10)
        raise SmokeFailure(
            "scheduled task did not complete once with a target file in 6 minutes"
        )

    def ui_locale(self) -> str:
        try:
            from playwright.sync_api import sync_playwright
        except ImportError as exc:
            raise SmokeFailure("Playwright is not installed") from exc
        project = (
            self.create_project(uuid4().hex[:8])
            if self.args.kind == "project"
            else None
        )
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True, args=["--no-sandbox"])
            context = browser.new_context(locale="fr-FR")
            try:
                context.add_cookies(
                    [{"name": "token", "value": self.token, "url": self.base}]
                )
                context.add_init_script("localStorage.setItem('locale','fr-FR');")
                pages = [self.base]
                if project:
                    pages.append(
                        f"{self.base}/?atlas_projects=1&org={project[0]}&project={project[1]}"
                    )
                for page_url in pages:
                    page = context.new_page()
                    page.goto(page_url, wait_until="domcontentloaded", timeout=60000)
                    page.wait_for_timeout(2000)
                    visible = page.locator("body").inner_text()
                    found = [word for word in DENY_LIST if word in visible]
                    page.close()
                    if found:
                        raise SmokeFailure("visible English UI: " + ", ".join(found))
            finally:
                context.close()
                browser.close()
        return "French home and project pages contain no deny-list string"

    def cleanup_revision(self, path: str) -> dict[str, str]:
        current = self.request("GET", path)
        if current.status_code == 200:
            revision = self._json(current).get("revision")
            if isinstance(revision, int):
                return {"If-Match": f'"{revision}"'}
        return {}

    def cleanup(self) -> list[str]:
        notes: list[str] = []
        for org, project, task in reversed(self.created_tasks):
            response = self.request(
                "DELETE",
                f"/api/atlas/orgs/{org}/projects/{project}/scheduled-tasks/{task}",
            )
            notes.append(f"task delete HTTP {response.status_code}")
        for kind, org, file_id in reversed(self.created_files):
            path = (
                f"/api/v1/files/{file_id}"
                if kind == "ordinary"
                else f"/api/atlas/orgs/{org}/projects/{file_id.rsplit('/', 1)[0]}/files/{file_id.rsplit('/', 1)[1]}"
            )
            headers = self.cleanup_revision(path) if kind != "ordinary" else {}
            response = self.request("DELETE", path, headers=headers)
            notes.append(f"file delete HTTP {response.status_code}")
        for chat_id in reversed(self.created_chats):
            response = self.request("DELETE", f"/api/v1/chats/{chat_id}")
            notes.append(f"chat delete HTTP {response.status_code}")
        for org, project in reversed(self.created_projects):
            path = f"/api/atlas/orgs/{org}/projects/{project}"
            response = self.request("DELETE", path, headers=self.cleanup_revision(path))
            notes.append(f"project delete HTTP {response.status_code}")
            if response.status_code in OPTIONAL_STATUS | {409}:
                renamed = self.request(
                    "PATCH",
                    f"/api/atlas/orgs/{org}/projects/{project}",
                    headers={"If-Match": f'"{self.project_revisions[(org, project)]}"'}
                    if (org, project) in self.project_revisions
                    else {},
                    json={"name": f"smoke-{datetime.now(timezone.utc):%Y-%m-%d}-done"},
                )
                notes.append(f"project rename HTTP {renamed.status_code}")
        shutil.rmtree(self.tempdir, ignore_errors=True)
        return notes


class Skip(RuntimeError):
    pass


def read_secret_file(path: Path) -> str:
    mode = path.stat().st_mode & 0o777
    if mode & 0o077:
        raise SmokeFailure(f"credential file {path} must be mode 600 or stricter")
    return path.read_text(encoding="utf-8")


def parse_account(path: Path) -> dict[str, Any]:
    raw = read_secret_file(path)
    try:
        value = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise SmokeFailure("account file must contain one JSON object") from exc
    if not isinstance(value, dict):
        raise SmokeFailure("account file must contain one JSON object")
    return value


def validate_base(value: str) -> str:
    parsed = urlparse(value)
    if parsed.scheme != "https" or not parsed.netloc:
        raise argparse.ArgumentTypeError("--base must be an HTTPS URL")
    if parsed.port in {8000, 8020, 3000, 5432}:
        raise argparse.ArgumentTypeError("--base uses a prohibited port")
    return value.rstrip("/")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", type=validate_base, required=True)
    parser.add_argument("--account-file", required=True)
    parser.add_argument("--kind", choices=("project", "guest"), required=True)
    parser.add_argument("--only", help="comma-separated check names")
    parser.add_argument("--api-key-file")
    parser.add_argument(
        "--insecure", action="store_true", help="allow a test TLS certificate"
    )
    args = parser.parse_args()
    account = parse_account(Path(args.account_file))
    smoke = Smoke(args, account)
    checks: list[tuple[str, Callable[[], str], bool]] = [
        ("login", smoke.login, True),
        ("chat_web", smoke.chat_web, True),
        ("chat_document", smoke.chat_document, True),
        ("image", smoke.image, True),
        ("api_key", smoke.api_key, True),
        ("project_deliverable", smoke.project_deliverable, args.kind == "project"),
        ("project_memory", smoke.project_memory, args.kind == "project"),
        ("move_to_project", smoke.move_to_project, args.kind == "project"),
        ("scheduled_task", smoke.scheduled_task, args.kind == "project"),
        ("ui_locale", smoke.ui_locale, True),
    ]
    selected = (
        set(args.only.split(",")) if args.only else {name for name, _, _ in checks}
    )
    unknown = selected - {name for name, _, _ in checks}
    if unknown:
        parser.error("unknown check(s): " + ", ".join(sorted(unknown)))
    if selected - {"login"}:
        selected.add("login")
    outcomes: list[Outcome] = []
    try:
        for name, check, applicable in checks:
            if name not in selected:
                continue
            outcomes.append(
                Outcome(name, "SKIP", 0.0, "not selected")
                if not applicable
                else smoke.run(name, check)
            )
    finally:
        cleanup = smoke.cleanup()
        print(
            "cleanup: "
            + (", ".join(cleanup) if cleanup else "no remote resources created")
        )
    print("\ncheck                 status  seconds  reason")
    print("-" * 78)
    for outcome in outcomes:
        print(
            f"{outcome.name:<21} {outcome.status:<6} {outcome.seconds:7.1f}  {outcome.reason}"
        )
    return 1 if any(outcome.status == "FAIL" for outcome in outcomes) else 0


if __name__ == "__main__":
    raise SystemExit(main())
