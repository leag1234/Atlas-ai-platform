"""Owner 2026-10-28 release defects, reproduced before the fifth-pass repairs."""

import hashlib
import io
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch
from uuid import uuid4

from openpyxl import load_workbook
from packages.doc_render import render
from packages.doc_render.schema import Content
from services.orchestrator.project_document_input import normalize
from services.orchestrator.project_document_tools import execute


def expenses() -> Content:
    return Content.model_validate(
        normalize(
            dict(
                path="totals.xlsx",
                formulas=True,
                content='# Costs\n## Expenses\n| Category | Item | Amount |\n| --- | --- | --- |\n| Formation | A | 200 |\n| Formation | B | 300 |\n| Equipement | C | 600 |\n| Equipement | D | 400 |\n| Formation total | — | =SUMIF(A2:A5,"Formation",C2:C5) |\n| Equipement total | — | =SUMIF(A2:A5,"Equipement",C2:C5) |\n| Grand total | — | =SUM(C2:C5) |',
            )
        )["content"]
    )


class SpreadsheetRowsTests(unittest.TestCase):
    def test_live_wrong_ranges_bind_to_rows_actually_written(self) -> None:
        for paragraphs in ((), ("Context before data",)):
            value = expenses().model_dump()
            value["sections"][0]["paragraphs"] = paragraphs
            with (
                self.subTest(paragraphs=paragraphs),
                tempfile.TemporaryDirectory() as directory,
            ):
                path = Path(directory) / "totals.xlsx"
                render(Content.model_validate(value), path, image_root=path.parent)
                book = load_workbook(path, data_only=True)
                formulas = load_workbook(path, data_only=False)
                try:
                    offset = len(paragraphs)
                    self.assertEqual(
                        [book["1"][f"C{r + offset}"].value for r in (8, 9, 10)],
                        [500, 1000, 1500],
                    )
                    self.assertEqual(
                        formulas["1"][f"C{8 + offset}"].value,
                        f'=SUMIF(A{4 + offset}:A{7 + offset},"Formation",C{4 + offset}:C{7 + offset})',
                    )
                finally:
                    book.close()
                    formulas.close()

    def test_cached_summary_must_match_selected_source_before_publication(self) -> None:
        from packages.doc_render.spreadsheet import validate_sources
        from packages.attachment_reader import read_structured

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "totals.xlsx"
            render(expenses(), path, image_root=path.parent)
            raw = path.read_bytes()
            source = read_structured(raw, ".xlsx", {"sheet": "1", "range": "A4:C7"})[
                "segments"
            ]
            validate_sources(raw, [source])
            changed = json.loads(json.dumps(source))
            next(s for s in changed if s["locator"]["cell"] == "C7")["text"] = "900"
            with self.assertRaisesRegex(ValueError, "spreadsheet_source_mismatch"):
                validate_sources(raw, [changed])
            # A tampered cached value must fail even with intact source rows.
            from zipfile import ZipFile, ZIP_DEFLATED
            from xml.etree import ElementTree as ET

            output = io.BytesIO()
            ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            with (
                ZipFile(io.BytesIO(raw)) as archive,
                ZipFile(output, "w", ZIP_DEFLATED) as target,
            ):
                for name in archive.namelist():
                    data = archive.read(name)
                    if name == "xl/worksheets/sheet1.xml":
                        tree = ET.fromstring(data)
                        for cell in tree.iter(ns + "c"):
                            if cell.attrib["r"] == "C9":
                                value = cell.find(ns + "v")
                                assert value is not None
                                value.text = "0"
                        data = ET.tostring(tree)
                    target.writestr(name, data)
            with self.assertRaisesRegex(ValueError, "spreadsheet_source_mismatch"):
                validate_sources(output.getvalue(), [source])

    def test_qualified_sheet_totals_are_checked_against_source(self) -> None:
        from packages.attachment_reader import read_structured
        from packages.doc_render.spreadsheet import validate_sources

        for start, end, accepted in ((4, 7, True), (2, 5, False)):
            value = expenses().model_dump()
            rows = [list(row) for row in value["sections"][0]["tables"][0]["rows"]]
            rows[5][2] = (
                f"=SUMIF('1'!A{start}:A{end},\"Equipement\",'1'!C{start}:C{end})"
            )
            value["sections"][0]["tables"][0]["rows"] = rows
            with (
                self.subTest(accepted=accepted),
                tempfile.TemporaryDirectory() as directory,
            ):
                path = Path(directory) / "totals.xlsx"
                render(Content.model_validate(value), path, image_root=path.parent)
                raw = path.read_bytes()
                source = read_structured(
                    raw, ".xlsx", {"sheet": "1", "range": "A4:C7"}
                )["segments"]
                if accepted:
                    validate_sources(raw, [source])
                else:
                    with self.assertRaisesRegex(
                        ValueError, "spreadsheet_source_mismatch"
                    ):
                        validate_sources(raw, [source])


class PublishedReadTests(unittest.IsolatedAsyncioTestCase):
    async def test_current_turn_receipt_allows_read_but_not_other_version(self) -> None:
        from services.retrieval.project_identity import ProjectError

        file_id, version, turn = (str(uuid4()) for _ in range(3))
        raw = b"published"
        execution = SimpleNamespace(
            request_id=turn,
            session={"base_file_versions": []},
            check=AsyncMock(),
            broker=SimpleNamespace(
                public=AsyncMock(
                    return_value=(
                        200,
                        dict(
                            version_id=version,
                            relative_path="totals.xlsx",
                            size_bytes=len(raw),
                            sha256=hashlib.sha256(raw).hexdigest(),
                        ),
                    )
                ),
                content=AsyncMock(return_value=raw),
            ),
        )
        tools = SimpleNamespace(
            execution=execution,
            deliverables=[dict(turn_id=turn, file_version_id=version)],
            workspace=SimpleNamespace(
                invoke=AsyncMock(return_value={"segments": [{"text": "1500"}]})
            ),
        )
        result = await execute(
            tools,
            "read_document",
            json.dumps(dict(file_id=file_id, version=version)),
            30,
        )
        self.assertEqual(result["segments"][0]["text"], "1500")
        for rejected in (str(uuid4()),):
            with self.assertRaises(ProjectError):
                await execute(
                    tools,
                    "read_document",
                    json.dumps(dict(file_id=file_id, version=rejected)),
                    30,
                )
        receipt_id = str(uuid4())
        tools.deliverables[0]["id"] = receipt_id
        metadata = execution.broker.public.return_value
        execution.broker.public.side_effect = [
            (200, {"items": [{"id": file_id, "version_id": version}]}),
            metadata,
        ]
        await execute(
            tools,
            "read_document",
            json.dumps(dict(file_id=receipt_id, version=version)),
            30,
        )
        self.assertEqual(tools.workspace.invoke.await_args.args[1]["file_id"], file_id)
        execution.broker.public.side_effect = None
        tools.deliverables[0]["turn_id"] = str(uuid4())
        with self.assertRaises(ProjectError):
            await execute(
                tools,
                "read_document",
                json.dumps(dict(file_id=file_id, version=version)),
                30,
            )


class PublicAnswerTests(unittest.TestCase):
    def test_memory_record_identifier_is_not_reply_content(self) -> None:
        from services.orchestrator.project_transcript import answer_content

        record = "df53e1cc-44c6-4a8a-a110-25b4bd8f90e1"
        text = (
            "Le budget maximum est de 2000 euros.\nSource : mémoire du projet (identifiant "
            + record
            + ")."
        )
        self.assertNotIn(record, answer_content(text))
        self.assertIn("2000 euros", answer_content(text))
        inline = "Le budget est de 2000 euros (mémoire ID " + record + ")."
        self.assertEqual(
            answer_content(json.dumps(dict(answer=inline, facts=[]))),
            "Le budget est de 2000 euros.",
        )

    def test_live_source_chunk_identifier_is_not_reply_content(self) -> None:
        from services.orchestrator.project_transcript import answer_content

        chunk = "d8bd5878272ab92f67870d89f059e2cef834dab75afafb9ffdb1b766edb726b1"
        text = "Totaux du fichier source (chunk `" + chunk + "` : Formation = 500)."
        self.assertEqual(
            answer_content(text), "Totaux du fichier source (Formation = 500)."
        )

    def test_live_forge_pr_url_survives_provenance_filter(self) -> None:
        from services.orchestrator.source_truth import (
            tool_result_urls,
            without_unretrieved,
        )

        url = "https://atlas-project-forge:3443/code-demo/demo/pulls/4"
        receipt = {"state": "succeeded", "pr_url": url}
        answer = "PR : [Voir la PR](" + url + ")"
        self.assertEqual(without_unretrieved(answer, tool_result_urls(receipt)), answer)
        invented = "https://unverified.example/pulls/99"
        self.assertNotIn(
            invented,
            without_unretrieved(
                answer + " [Other](" + invented + ")", tool_result_urls(receipt)
            ),
        )
        self.assertNotIn(url, tool_result_urls({"untrusted_text": url}))

    def test_standalone_route_enters_shared_shell_and_preserves_scope(self) -> None:
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from services.orchestrator.project_ui import router

        app = FastAPI()
        app.include_router(router)
        with (
            patch.dict("os.environ", {"ATLAS_PROJECT_MODE": "1"}),
            TestClient(app) as client,
        ):
            response = client.get(
                "/project-mode?org=one&project=two&conversation=three",
                follow_redirects=False,
            )
        self.assertEqual(response.status_code, 307)
        self.assertEqual(
            response.headers["location"],
            "/?atlas_projects=1&org=one&project=two&conversation=three",
        )


class ForgeBindingTests(unittest.TestCase):
    def test_scratch_target_and_sha_bind_to_remote_canonical_base(self) -> None:
        from unittest.mock import Mock
        from services.retrieval.project_work_forge import bind_publication
        from services.orchestrator.project_forge import GitHubForge, ForgeError

        adapter = Mock()
        adapter.request.return_value = {"default_branch": "main"}
        adapter.refs.return_value = {"main": "b" * 40, "previous-proposal": "b" * 40}
        adapter.connector_digest.return_value = "connector"
        action = dict(
            kind="publish_change",
            base_commit="a" * 40,
            parameters=dict(
                branch="previous-proposal",
                target="master",
                message="Average",
                title="Average",
                body="Tests passed",
            ),
        )
        with self.assertRaisesRegex(ForgeError, "base_mismatch"):
            GitHubForge.bind(adapter, action)
        bound = bind_publication({"base_commit": "b" * 40}, action, adapter)
        self.assertEqual(bound["base_commit"], "b" * 40)
        self.assertEqual(bound["parameters"]["target"], "main")
        self.assertNotIn(bound["parameters"]["branch"], adapter.refs.return_value)
        self.assertEqual(GitHubForge.bind(adapter, bound)["target_commit"], "b" * 40)
        adapter.refs.return_value["main"] = "c" * 40
        with self.assertRaises(ForgeError):
            bind_publication({"base_commit": "b" * 40}, action, adapter)


class PublicationSourceGuardTests(unittest.IsolatedAsyncioTestCase):
    async def test_wrong_source_total_never_reaches_publisher(self) -> None:
        from services.orchestrator.project_document_tools import publish
        from services.orchestrator.terminal_client import (
            VerifiedDocument,
            ProjectPublisher,
        )
        from packages.attachment_reader import read_structured
        from unittest.mock import Mock

        file_id, version = str(uuid4()), str(uuid4())
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "totals.xlsx"
            render(expenses(), path, image_root=path.parent)
            raw = path.read_bytes()
        segments = read_structured(raw, ".xlsx", {"sheet": "1", "range": "A4:C7"})[
            "segments"
        ]
        next(s for s in segments if s["locator"]["cell"] == "C7")["text"] = "900"
        execution = Mock()
        execution.check = AsyncMock()
        execution.session = {
            "base_file_versions": [dict(file_id=file_id, version_id=version)]
        }
        execution.broker.public = AsyncMock(
            side_effect=[
                (200, {"deliverables_folder_id": "folder"}),
                (200, {"items": [{"id": "folder", "revision": 0}]}),
                (
                    200,
                    dict(
                        version_id=version,
                        relative_path="source.xlsx",
                        size_bytes=len(raw),
                        sha256=hashlib.sha256(raw).hexdigest(),
                    ),
                ),
            ]
        )
        execution.broker.content = AsyncMock(return_value=raw)
        tools = Mock(execution=execution)
        tools.workspace.invoke = AsyncMock(
            return_value={"coverage": {"complete": True}, "segments": segments}
        )
        document = VerifiedDocument(raw, b"preview", hashlib.sha256(raw).hexdigest())
        with (
            patch(
                "services.orchestrator.presentation_delivery.build_project_document",
                return_value=document,
            ),
            patch.object(ProjectPublisher, "publish", new_callable=AsyncMock) as write,
        ):
            with self.assertRaisesRegex(ValueError, "spreadsheet_source_mismatch"):
                await publish(
                    tools,
                    dict(
                        path="totals.xlsx",
                        content=expenses().model_dump(),
                        sources=[dict(file_id=file_id, version=version)],
                        strategy="created",
                    ),
                    30,
                )
            write.assert_not_awaited()
