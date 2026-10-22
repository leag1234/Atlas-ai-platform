"""Offline regression coverage for deterministic filesystem attachment reads."""

import asyncio
import base64
from decimal import Decimal
import io
import json
import os
from pathlib import Path
import subprocess
import tempfile
import time
import unittest
from unittest.mock import AsyncMock, patch
import zipfile

from docx import Document
from pypdf import PdfWriter
from packages.language import conversation_language
from services.orchestrator.chat_schema import ChatMessage, ChatRequest
from services.orchestrator.chat_stream import response
from services.orchestrator.interactions import Interaction
from services.orchestrator.model import Configuration, GatewayModel
from services.orchestrator.loop import Reservation, Turn
from services.orchestrator.terminal_client import TerminalClient
from test_m23_documents import pdf

FIXTURE = json.loads(Path("tests/fixtures/attachment-language.json").read_text())


def announced(paths: list[str], question: str) -> str:
    from html import escape

    return (
        "<attached_files>\n"
        + "\n".join(
            '<file type="filesystem" name="Annual English report" url="'
            + escape(path, quote=True)
            + '" />'
            for path in paths
        )
        + "\n</attached_files>\n"
        + question
    )


def office(path: Path, entries: dict[str, str]) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)


class AttachmentLanguageTests(unittest.TestCase):
    def test_markup_does_not_vote_for_answer_language(self) -> None:
        for key, expected in (("fr", "fr"), ("short_fr", "fr"), ("en", "en")):
            message = ChatMessage(
                role="user",
                content=announced(
                    ["/home/user/report.pdf", "/home/user/notes.docx"], FIXTURE[key]
                ),
            )
            self.assertEqual(conversation_language([message], "en"), expected)


class AttachmentExtractionTests(unittest.TestCase):
    def test_formats_and_page_sheet_slide_markers(self) -> None:
        from packages.attachment_reader import read_attachment

        with tempfile.TemporaryDirectory(dir="BRAIN") as root:
            directory = Path(root)
            (directory / "sample.pdf").write_bytes(pdf(2))
            document = Document()
            document.add_paragraph(FIXTURE["fact"])
            document.add_table(1, 1).cell(0, 0).text = "Table finding"
            document.save(str(directory / "sample.docx"))
            office(
                directory / "sample.xlsx",
                {
                    "xl/workbook.xml": '<workbook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Revenue" r:id="rId1"/></sheets></workbook>',
                    "xl/_rels/workbook.xml.rels": '<Relationships><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>',
                    "xl/sharedStrings.xml": "<sst><si><t>Revenue total</t></si></sst>",
                    "xl/worksheets/sheet1.xml": '<worksheet><sheetData><row><c r="A1" t="s"><v>0</v></c><c r="B1"><v>730</v></c></row></sheetData></worksheet>',
                },
            )
            office(
                directory / "sample.pptx",
                {
                    "ppt/presentation.xml": '<presentation xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sldIdLst><sldId r:id="rId1"/></sldIdLst></presentation>',
                    "ppt/_rels/presentation.xml.rels": '<Relationships><Relationship Id="rId1" Target="slides/slide1.xml"/></Relationships>',
                    "ppt/slides/slide1.xml": "<slide><p><t>Final recommendation</t></p></slide>",
                },
            )
            (directory / "sample.txt").write_text(FIXTURE["fact"])
            expected = {
                "pdf": ("[Page 2]", "730 euros"),
                "docx": (FIXTURE["fact"], "Table finding"),
                "xlsx": ("[Sheet: Revenue]", "730"),
                "pptx": ("[Slide 1]", "Final recommendation"),
                "txt": (FIXTURE["fact"],),
            }
            for extension, facts in expected.items():
                result = read_attachment(
                    directory / ("sample." + extension), "findings"
                )
                self.assertFalse(result["errors"])
                for fact in facts:
                    self.assertIn(fact, str(result["text"]))

    def test_bad_file_scanned_page_and_long_document_are_honest(self) -> None:
        from packages.attachment_reader import read_attachment

        with tempfile.TemporaryDirectory(dir="BRAIN") as root:
            directory = Path(root)
            bad = directory / "bad.pdf"
            bad.write_bytes(b"invalid")
            result = read_attachment(bad, "summary")
            self.assertIn("PdfReadError", str(result["errors"]))
            scanned = directory / "scan.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            buffer = io.BytesIO()
            writer.write(buffer)
            scanned.write_bytes(buffer.getvalue())
            self.assertIn(
                "no text layer", str(read_attachment(scanned, "summary")["errors"])
            )
            long = directory / "long.pdf"
            long.write_bytes(pdf(30))
            result = read_attachment(long, "Initial budget and quality review")
            self.assertTrue(result["hierarchical_synthesis"])
            self.assertLess(len(str(result["text"]).encode()), 12000)
            self.assertEqual(result["pages"], 30)
            self.assertIn("730", str(result["text"]))
            self.assertIn(
                "FileNotFoundError",
                str(read_attachment(directory / "absent.txt", "summary")["errors"]),
            )

    def test_presentation_text_tables_notes_and_template_are_read_in_order(
        self,
    ) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        from packages.attachment_reader import read_attachment

        with tempfile.TemporaryDirectory(dir="BRAIN") as root:
            directory = Path(root)
            deck = Presentation()
            first = deck.slides.add_slide(deck.slide_layouts[1])
            first.shapes.title.text = "First title"
            first.placeholders[1].text = "First body"
            table = first.shapes.add_table(
                2, 2, Inches(1), Inches(3), Inches(4), Inches(1)
            ).table
            table.cell(0, 0).text = "Metric"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "Revenue"
            table.cell(1, 1).text = "730"
            first.notes_slide.notes_text_frame.text = "First speaker note"
            second = deck.slides.add_slide(deck.slide_layouts[1])
            second.shapes.title.text = "Second title"
            second.placeholders[1].text = "Second body"
            second.notes_slide.notes_text_frame.text = "Second speaker note"
            pptx = directory / "deck.pptx"
            deck.save(str(pptx))
            template = directory / "deck.potx"
            template.write_bytes(pptx.read_bytes())

            for path in (pptx, template):
                result = read_attachment(path, "summarize")
                self.assertFalse(result["errors"])
                self.assertEqual(result["pages"], 2)
                text = str(result["text"])
                self.assertIn("[Slide 1]", text)
                self.assertIn("First title", text)
                self.assertIn("Revenue", text)
                self.assertIn("730", text)
                self.assertIn("[Speaker notes]", text)
                self.assertIn("First speaker note", text)
                self.assertLess(text.index("First title"), text.index("Second title"))
                self.assertLess(
                    text.index("First speaker note"), text.index("Second speaker note")
                )

            from services.orchestrator.documents import Attachment, extract_document

            extracted = extract_document(
                Attachment(
                    filename="deck.pptx",
                    file_data=base64.b64encode(pptx.read_bytes()).decode(),
                )
            )
            self.assertEqual(extracted.pages, 2)
            self.assertIn("Second speaker note", extracted.text)


class StreamExceptionTests(unittest.IsolatedAsyncioTestCase):
    async def test_unexpected_exception_is_diagnosable_without_input_or_cost(
        self,
    ) -> None:
        async def broken(request: ChatRequest, item: Interaction) -> None:
            item.cout_eur = 0.30
            raise LookupError("sensitive-request-and-document-content")

        request = ChatRequest(messages=[ChatMessage(role="user", content="Summarize")])
        with (
            tempfile.TemporaryDirectory(dir="BRAIN") as root,
            patch.dict(os.environ, {"ATLAS_INTERACTION_DIR": root}),
            self.assertLogs("services.orchestrator.chat_stream", level="ERROR") as logs,
        ):
            events = [
                str(event)
                async for event in response(
                    request, Interaction(), time.monotonic(), broken
                ).body_iterator
            ]
        public = "".join(events)
        self.assertIn("unexpected failure", public.lower())
        self.assertIn("not a budget or deadline refusal", public.lower())
        self.assertIn("LookupError", "".join(logs.output))
        self.assertIn("broken", "".join(logs.output))
        self.assertNotIn(
            "sensitive-request-and-document-content", public + "".join(logs.output)
        )


class AttachmentPipelineTests(unittest.IsolatedAsyncioTestCase):
    async def test_actual_helper_runs_before_model_and_partial_failure_reaches_user(
        self,
    ) -> None:
        from services.orchestrator.chat_pipeline import process

        with tempfile.TemporaryDirectory(dir="BRAIN") as root:
            directory = Path(root).resolve()
            good = directory / "notes.txt"
            good.write_text(FIXTURE["fact"])
            bad = directory / "broken.pdf"
            bad.write_bytes(b"invalid")
            request = ChatRequest(
                messages=[
                    ChatMessage(
                        role="user",
                        content=announced([str(bad), str(good)], FIXTURE["fr"]),
                    )
                ],
                ui_locale="en",
            )
            model = GatewayModel(
                "http://unused",
                Configuration(
                    input_eur_per_mtok=Decimal(0),
                    output_eur_per_mtok=Decimal(0),
                    max_tokens=100,
                ),
            )

            async def complete(
                messages: list[dict[str, object]], timeout: float
            ) -> Turn:
                text = json.dumps(messages, ensure_ascii=False)
                self.assertIn(FIXTURE["fact"], text)
                self.assertIn("PdfReadError", text)
                self.assertIn("Answer in French.", text)
                return Turn(FIXTURE["fact"], usage=Reservation(1, Decimal(0)))

            async def execute(
                client: TerminalClient, command: str, timeout: float
            ) -> dict[str, object]:
                client.commands += 1
                result = await asyncio.to_thread(
                    subprocess.run,
                    command,
                    shell=True,
                    capture_output=True,
                    text=True,
                    timeout=timeout,
                    cwd=directory,
                )
                self.assertEqual(result.returncode, 0, result.stderr)
                return {
                    "exit_code": 0,
                    "output": [{"type": "output", "data": result.stdout}],
                }

            with (
                patch.dict(
                    os.environ,
                    {
                        "ATLAS_TERMINAL_ENABLED": "1",
                        "ATLAS_WEB_CACHE": str(directory / "cache"),
                        "GPU_LOCAL": "0",
                        "ATLAS_SEARCH_PROVIDER": "serpapi",
                    },
                ),
                patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
                patch.object(model, "complete", side_effect=complete),
                patch.object(TerminalClient, "initialize", AsyncMock()),
                patch.object(TerminalClient, "execute", execute),
            ):
                item = Interaction()
                await process(request, item)
            self.assertEqual(item.state, "done")
            self.assertIn(FIXTURE["fact"], item.reponse)
            self.assertIn("broken.pdf", item.reponse)
            self.assertIn("PdfReadError", item.reponse)
            self.assertNotIn("alternatives", item.reponse)
            self.assertEqual(len(item.documents), 2)
            self.assertGreater(item.terminal_commands, 0)
            self.assertTrue(model.has_attachments)


class DirectAttachmentTests(unittest.IsolatedAsyncioTestCase):
    async def test_mixed_documents_and_images_use_both_paths(self) -> None:
        """Regression: direct uploads must not drop images when a document is present."""
        import base64
        from services.orchestrator.documents import Attachment, Extracted
        from test_vision_schema import picture

        from services.orchestrator.chat_pipeline import process

        model = GatewayModel(
            "http://unused",
            Configuration(
                input_eur_per_mtok=Decimal(0),
                output_eur_per_mtok=Decimal(0),
                max_tokens=100,
            ),
        )

        async def complete(messages: list[dict[str, object]], timeout: float) -> Turn:
            evidence = json.dumps(messages, ensure_ascii=False)
            self.assertIn("SCREENSHOT FACT", evidence)
            self.assertIn("document fact", evidence)
            return Turn(
                "Image: SCREENSHOT FACT. Document: document fact.",
                usage=Reservation(1, Decimal("0.01")),
            )

        async def extract(
            attachment: Attachment, **_: object
        ) -> tuple[Extracted, Decimal, list[int]]:
            return (
                Extracted(attachment.filename, 1, 1, "document fact"),
                Decimal(0),
                [],
            )

        vision = {
            "text": "SCREENSHOT FACT: readable synthetic screenshot.",
            "usage": {"prompt_tokens": 10, "completion_tokens": 5},
            "cost_eur": "0.01",
            "observation": {"provider": "escalade", "route": "complexe"},
        }
        for suffix in ("pdf", "docx", "xlsx", "pptx"):
            with self.subTest(suffix=suffix):
                request = ChatRequest(
                    messages=[
                        ChatMessage.model_validate(
                            {
                                "role": "user",
                                "content": [
                                    {
                                        "type": "text",
                                        "text": "Describe image and summarize document",
                                    },
                                    {
                                        "type": "image_url",
                                        "image_url": {"url": picture()},
                                    },
                                    {
                                        "type": "image_url",
                                        "image_url": {"url": picture()},
                                    },
                                ],
                            }
                        )
                    ],
                    documents=[
                        Attachment(
                            filename="source." + suffix,
                            file_data=base64.b64encode(b"placeholder").decode(),
                        )
                    ],
                )
                with (
                    patch.dict(os.environ, {"ATLAS_TERMINAL_ENABLED": "1"}),
                    patch.object(
                        GatewayModel, "connect", AsyncMock(return_value=model)
                    ),
                    patch.object(
                        GatewayModel, "post", AsyncMock(return_value=vision)
                    ) as post,
                    patch.object(model, "complete", side_effect=complete),
                    patch(
                        "services.orchestrator.attachment_reading.extract_attachment_with_ocr",
                        new=AsyncMock(side_effect=extract),
                    ),
                ):
                    item = Interaction()
                    await process(request, item)
                self.assertEqual(item.state, "done")
                self.assertIn("SCREENSHOT FACT", item.reponse)
                self.assertIn("document fact", item.reponse)
                self.assertIsNotNone(post.await_args)
                assert post.await_args is not None
                vision_parts = post.await_args.args[1]["messages"][0]["content"]
                self.assertEqual(
                    sum(part["type"] == "image_url" for part in vision_parts), 2
                )
                self.assertEqual(len(item.documents), 1)

    async def test_mixed_attachment_keeps_document_when_vision_fails(self) -> None:
        import base64
        from services.orchestrator.documents import Attachment, Extracted
        from test_vision_schema import picture

        from services.orchestrator.chat_pipeline import process
        from services.orchestrator.model import GatewayError

        request = ChatRequest(
            messages=[
                ChatMessage.model_validate(
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": "décris l'image et résume le document",
                            },
                            {"type": "image_url", "image_url": {"url": picture()}},
                        ],
                    }
                )
            ],
            documents=[
                Attachment(
                    filename="source.pdf", file_data=base64.b64encode(b"x").decode()
                )
            ],
        )
        model = GatewayModel(
            "http://unused",
            Configuration(
                input_eur_per_mtok=Decimal(0),
                output_eur_per_mtok=Decimal(0),
                max_tokens=100,
            ),
        )
        with (
            patch.dict(os.environ, {"ATLAS_TERMINAL_ENABLED": "0"}),
            patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
            patch.object(
                GatewayModel,
                "post",
                AsyncMock(side_effect=GatewayError("provider_error", 502)),
            ),
            patch.object(
                model,
                "complete",
                AsyncMock(
                    return_value=Turn(
                        "Résumé du document.", usage=Reservation(1, Decimal("0.01"))
                    )
                ),
            ),
            patch(
                "services.orchestrator.attachment_reading.extract_attachment_with_ocr",
                new=AsyncMock(
                    return_value=(
                        Extracted("source.pdf", 1, 1, "document fact"),
                        Decimal(0),
                        [],
                    )
                ),
            ),
        ):
            item = Interaction()
            await process(request, item)
        self.assertIn("Résumé du document.", item.reponse)
        self.assertIn("Je n’ai pas pu voir l’image", item.reponse)
        self.assertIn("provider_error", item.reponse)

    async def test_direct_documents_use_french_and_keep_readable_files(self) -> None:
        import base64
        from services.orchestrator.chat_pipeline import process
        from services.orchestrator.documents import Attachment

        request = ChatRequest(
            ui_locale="en",
            messages=[ChatMessage(role="user", content=FIXTURE["fr"])],
            documents=[
                Attachment(
                    filename="broken.pdf",
                    file_data=base64.b64encode(b"invalid").decode(),
                ),
                Attachment(
                    filename="notes.txt",
                    file_data=base64.b64encode(FIXTURE["fact"].encode()).decode(),
                ),
            ],
        )
        model = GatewayModel(
            "http://unused",
            Configuration(
                input_eur_per_mtok=Decimal(0),
                output_eur_per_mtok=Decimal(0),
                max_tokens=100,
            ),
        )

        async def complete(messages: list[dict[str, object]], timeout: float) -> Turn:
            text = json.dumps(messages, ensure_ascii=False)
            self.assertIn("Answer in French.", text)
            self.assertIn(FIXTURE["fact"], text)
            return Turn(FIXTURE["fact"], usage=Reservation(1, Decimal(0)))

        with (
            patch.dict(os.environ, {"ATLAS_TERMINAL_ENABLED": "0"}),
            patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
            patch.object(model, "complete", side_effect=complete),
        ):
            item = Interaction()
            await process(request, item)
        self.assertIn("broken.pdf", item.reponse)
        self.assertIn(FIXTURE["fact"], item.reponse)
        self.assertEqual(item.state, "done")


class InlineSandboxTests(unittest.IsolatedAsyncioTestCase):
    async def test_uploaded_bytes_are_read_before_model_without_script_choice(
        self,
    ) -> None:
        from services.orchestrator.chat_pipeline import process
        from services.orchestrator.documents import Attachment
        import base64

        request = ChatRequest(
            messages=[ChatMessage(role="user", content=FIXTURE["en"])],
            documents=[
                Attachment(
                    filename="notes.txt",
                    file_data=base64.b64encode(b"Actual source finding 730").decode(),
                ),
                Attachment(
                    filename="more.txt",
                    file_data=base64.b64encode(b"Second source finding 19").decode(),
                ),
            ],
        )
        model = GatewayModel(
            "http://unused",
            Configuration(
                input_eur_per_mtok=Decimal(0),
                output_eur_per_mtok=Decimal(0),
                max_tokens=100,
            ),
        )

        async def complete(messages: list[dict[str, object]], timeout: float) -> Turn:
            text = json.dumps(messages)
            self.assertIn("Actual source finding 730", text)
            self.assertIn("Second source finding 19", text)
            self.assertNotIn("terminal_command", json.dumps(model.tools))
            return Turn("730 and 19", usage=Reservation(1, Decimal(0)))

        with (
            tempfile.TemporaryDirectory(dir="BRAIN") as root,
            patch.dict(
                os.environ,
                {
                    "ATLAS_TERMINAL_ENABLED": "1",
                    "ATLAS_WEB_CACHE": root + "/cache",
                    "GPU_LOCAL": "0",
                    "ATLAS_SEARCH_PROVIDER": "serpapi",
                },
            ),
            patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
            patch.object(model, "complete", side_effect=complete),
            patch.object(TerminalClient, "initialize", AsyncMock()),
            patch.object(
                TerminalClient,
                "upload",
                AsyncMock(return_value="/home/user/requests/test/source.txt"),
            ),
            patch.object(
                TerminalClient,
                "execute",
                AsyncMock(side_effect=AssertionError("Model must not choose a reader")),
            ),
        ):
            item = Interaction()
            await process(request, item)
        self.assertEqual(item.reponse, "730 and 19")
        self.assertEqual(len(item.documents), 2)


class ReaderBoundaryTests(unittest.TestCase):
    def test_partial_scanned_pdf_keeps_text_and_names_unread_page(self) -> None:
        import base64
        from packages.attachment_reader import read_attachment
        from services.orchestrator.documents import Attachment, extract_document

        writer = PdfWriter()
        writer.append(io.BytesIO(pdf(1)))
        writer.add_blank_page(width=100, height=100)
        writer.append(io.BytesIO(pdf(1)))
        buffer = io.BytesIO()
        writer.write(buffer)
        with tempfile.TemporaryDirectory(dir="BRAIN") as root:
            path = Path(root) / "mixed.pdf"
            path.write_bytes(buffer.getvalue())
            result = read_attachment(path, "budget")
        self.assertIn("[Page 3]", str(result["text"]))
        self.assertIn("page 2/3", str(result["errors"]))
        document = extract_document(
            Attachment(
                filename="mixed.pdf",
                file_data=base64.b64encode(buffer.getvalue()).decode(),
            )
        )
        self.assertIn("[Page 3]", document.text)
        self.assertIn("page 2/3", str(document.errors))

    def test_helper_quotes_paths_and_refuses_symlink_escape(self) -> None:
        from services.orchestrator.attachment_reading import (
            reader_command,
            decode_output,
        )

        with tempfile.TemporaryDirectory(dir="BRAIN") as root:
            directory = Path(root).resolve()
            source = directory / "notes'; $(touch injected).txt"
            source.write_text("Actual safe content")
            for path, expected in (
                (source, "Actual safe content"),
                (directory / "escape.txt", "outside the permitted upload area"),
            ):
                if path.name == "escape.txt":
                    path.symlink_to("/etc/hostname")
                result = subprocess.run(
                    reader_command(str(path), "summary"),
                    shell=True,
                    cwd=directory,
                    text=True,
                    capture_output=True,
                    check=True,
                )
                decoded = decode_output(
                    {
                        "exit_code": 0,
                        "output": [{"type": "output", "data": result.stdout}],
                    }
                )
                self.assertIn(expected, str(decoded))
            self.assertFalse((directory / "injected").exists())

    def test_reader_output_rejects_truncation_and_malformed_results(self) -> None:
        from services.orchestrator.attachment_reading import decode_output

        result = decode_output({"error": "terminal_output_truncated"})
        self.assertEqual(result["text"], "")
        self.assertIn("terminal_output_truncated", str(result["errors"]))
        with self.assertRaises(ValueError):
            decode_output(
                {"exit_code": 0, "output": [{"type": "output", "data": "{}"}]}
            )


class UnavailableAttachmentTests(unittest.IsolatedAsyncioTestCase):
    async def test_sandbox_setup_failure_names_both_files_and_actual_cause(
        self,
    ) -> None:
        from services.orchestrator.chat_pipeline import process

        request = ChatRequest(
            ui_locale="en",
            messages=[
                ChatMessage(
                    role="user",
                    content=announced(
                        ["/home/user/one.pdf", "/home/user/two.docx"], FIXTURE["en"]
                    ),
                )
            ],
        )
        model = GatewayModel(
            "http://unused",
            Configuration(
                input_eur_per_mtok=Decimal(0),
                output_eur_per_mtok=Decimal(0),
                max_tokens=100,
            ),
        )
        with (
            tempfile.TemporaryDirectory(dir="BRAIN") as root,
            patch.dict(
                os.environ,
                {
                    "ATLAS_TERMINAL_ENABLED": "1",
                    "ATLAS_WEB_CACHE": root + "/cache",
                    "ATLAS_SEARCH_PROVIDER": "serpapi",
                    "GPU_LOCAL": "0",
                },
            ),
            patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
            patch.object(
                model,
                "complete",
                AsyncMock(side_effect=AssertionError("No readable evidence")),
            ) as complete,
            patch.object(
                TerminalClient,
                "initialize",
                AsyncMock(side_effect=RuntimeError("terminal_http_401")),
            ),
            patch.object(TerminalClient, "execute", AsyncMock()) as execute,
        ):
            item = Interaction()
            await process(request, item)
        self.assertEqual(item.state, "done")
        self.assertEqual(item.cout_eur, 0)
        self.assertIn("one.pdf", item.reponse)
        self.assertIn("two.docx", item.reponse)
        self.assertIn("terminal_http_401", item.reponse)
        self.assertNotIn("alternatives", item.reponse)
        complete.assert_not_awaited()
        execute.assert_not_awaited()


class DocumentStreamErrorTests(unittest.IsolatedAsyncioTestCase):
    async def test_document_generation_exception_keeps_internal_cause(self) -> None:
        import base64
        from services.orchestrator.chat_pipeline import process
        from services.orchestrator.documents import Attachment

        request = ChatRequest(
            messages=[ChatMessage(role="user", content=FIXTURE["en"])],
            documents=[
                Attachment(
                    filename="notes.txt",
                    file_data=base64.b64encode(b"A readable source").decode(),
                )
            ],
        )
        model = GatewayModel(
            "http://unused",
            Configuration(
                input_eur_per_mtok=Decimal(0),
                output_eur_per_mtok=Decimal(0),
                max_tokens=100,
            ),
        )
        with (
            tempfile.TemporaryDirectory(dir="BRAIN") as root,
            patch.dict(
                os.environ,
                {"ATLAS_TERMINAL_ENABLED": "0", "ATLAS_INTERACTION_DIR": root},
            ),
            patch.object(GatewayModel, "connect", AsyncMock(return_value=model)),
            patch.object(
                model, "complete", AsyncMock(side_effect=ValueError("private-input"))
            ),
            self.assertLogs("services.orchestrator.chat_stream", level="ERROR") as logs,
        ):
            output = "".join(
                [
                    str(event)
                    async for event in response(
                        request, Interaction(), time.monotonic(), process
                    ).body_iterator
                ]
            )
        self.assertIn("unexpected failure", output.lower())
        self.assertIn("not a budget or deadline refusal", output.lower())
        self.assertIn("ValueError", "".join(logs.output))
        self.assertNotIn("private-input", output + "".join(logs.output))


def setUpModule() -> None:
    # BRAIN/ is ignored by git, so a clean checkout (CI) does not have it.
    import pathlib

    pathlib.Path("BRAIN").mkdir(exist_ok=True)
