"""REQ-PROJ-05: explicit formulas, computed caches and rejected external formulas."""

import tempfile
from pathlib import Path
import unittest

from openpyxl import load_workbook
from packages.doc_render import render
from packages.doc_render.schema import Content
from services.orchestrator.project_document_input import normalize


class FormulaDeliveryTests(unittest.TestCase):
    def content(self, formula: str = '=SUMIF(A4:A7,"Formation",B4:B7)') -> Content:
        value = normalize(
            dict(
                path="totals.xlsx",
                formulas=True,
                content="# Costs\n## Expenses\n| Category | Amount |\n| --- | --- |\n| Formation | 200 |\n| Formation | 300 |\n| Equipement | 600 |\n| Equipement | 400 |\n| Formation total | "
                + formula
                + ' |\n| Equipement total | =SUMIF(A4:A7,"Equipement",B4:B7) |\n| Grand total | =SUM(B8:B9) |',
            )
        )
        return Content.model_validate(value["content"])

    def test_real_formulas_and_calculated_values_reopen(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "totals.xlsx"
            render(self.content(), path, image_root=path.parent)
            formulas = load_workbook(path, data_only=False)
            values = load_workbook(path, data_only=True)
            try:
                self.assertEqual(formulas["1"]["B8"].data_type, "f")
                self.assertEqual(formulas["1"]["B4"].data_type, "n")
                self.assertEqual(
                    [values["1"][f"B{i}"].value for i in (8, 9, 10)], [500, 1000, 1500]
                )
            finally:
                formulas.close()
                values.close()

    def test_formula_publication_reopens_and_reports_only_safe_validation_reason(
        self,
    ) -> None:
        from services.orchestrator.presentation_delivery import build_project_document

        document = build_project_document(
            self.content(),
            ".xlsx",
            ["Formation", "200"],
            max_bytes=16777216,
            timeout=30,
        )
        self.assertGreater(len(document.content), 100)
        with self.assertRaisesRegex(ValueError, "unsupported_formula_function"):
            build_project_document(
                self.content("=UNKNOWN(B4:B7)"),
                ".xlsx",
                [],
                max_bytes=16777216,
                timeout=30,
            )

    def test_fractional_formula_cache_round_trips(self) -> None:
        value = self.content("=SUM(0.1,0.2)")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fraction.xlsx"
            render(value, path, image_root=path.parent)
            book = load_workbook(path, data_only=True)
            try:
                self.assertAlmostEqual(book["1"]["B8"].value, 0.3)
            finally:
                book.close()

    def test_untrusted_functions_external_references_and_cycles_fail_closed(
        self,
    ) -> None:
        for formula in (
            '=WEBSERVICE("https://example.org")',
            "='[other.xlsx]1'!A1",
            "=SUM(B8)",
            "=SUM(A1:A99999)",
        ):
            with (
                self.subTest(formula=formula),
                tempfile.TemporaryDirectory() as directory,
            ):
                path = Path(directory) / "totals.xlsx"
                with self.assertRaises(ValueError):
                    render(self.content(formula), path, image_root=path.parent)
                self.assertFalse(path.exists())

    def test_generic_workbooks_keep_formula_shaped_text_literal(self) -> None:
        value = self.content().model_dump()
        value["sections"][0]["tables"][0]["formulas"] = False
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "literal.xlsx"
            render(Content.model_validate(value), path, image_root=path.parent)
            book = load_workbook(path)
            try:
                self.assertEqual(book["1"]["B8"].data_type, "s")
            finally:
                book.close()

    def test_markdown_empty_cell_is_a_visible_placeholder_without_padding_rows(
        self,
    ) -> None:
        value = normalize(
            dict(
                path="summary.docx",
                content="| Name | Amount | Note |\n| --- | --- | --- |\n| Total | 1500 |  |",
            )
        )
        content = Content.model_validate(value["content"])
        self.assertEqual(content.sections[0].tables[0].rows[0], ("Total", "1500", "—"))
        with self.assertRaises(ValueError):
            normalize(
                dict(
                    path="summary.docx",
                    content="| Name | Amount |\n| --- | --- |\n| Total |",
                )
            )


class SourceAppendixTests(unittest.TestCase):
    def test_summary_keeps_cited_source_excerpts_in_verified_artifact(self) -> None:
        from services.orchestrator.project_document_input import markdown
        from services.orchestrator.project_document_tools import include_sources
        from services.orchestrator.presentation_delivery import build_project_document

        summary = markdown("# Synthèse Quartz\nRésumé du projet.", "Quartz")
        source = "Budget initial : 1500 euros.\nLivraison : décembre."
        content = include_sources(summary, [source])
        self.assertEqual(content.sections[0], summary.sections[0])
        self.assertEqual(include_sources(content, [source]), content)
        document = build_project_document(
            content, ".docx", [source], max_bytes=16777216, timeout=30
        )
        self.assertIn(b"Budget initial : 1500 euros.", document.preview)
