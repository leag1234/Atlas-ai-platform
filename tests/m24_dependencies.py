"""Owned, contained file dependencies for live recording; no provider probes."""

import asyncio
import base64
from contextlib import ExitStack, contextmanager
from collections.abc import Iterator
import hashlib
import io
import json
import os
from pathlib import Path
import secrets
import shlex
import socket
import subprocess
from unittest.mock import patch
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from packages.ports import atlas_ports
from services.orchestrator import terminal_stack
from services.orchestrator.terminal_client import TerminalClient
from m24_capture import download_files

FILE_CASES = {"J42", "J43", "J44", "formats_write"}


class DependencyError(RuntimeError):
    """Stage-only diagnostics: never include subprocess output or credentials."""


def preflight() -> None:
    """At most 60 s of terminal work plus two 10 s authenticated HTTP calls."""
    stage = "credentials"
    terminal = None
    error = None
    try:
        if not os.environ.get("OPEN_TERMINAL_API_KEY", "").strip():
            raise ValueError
        terminal = TerminalClient("chatcmpl-" + uuid4().hex)
        stage = "initialize"
        asyncio.run(asyncio.wait_for(terminal.initialize(), 5))
        stage = "upload"
        workbook = Workbook()
        workbook.worksheets[0].append(["amount", 7])
        buffer = io.BytesIO()
        workbook.save(buffer)
        path = asyncio.run(
            asyncio.wait_for(
                terminal.upload_generated("preflight.xlsx", buffer.getvalue(), 10), 10
            )
        )
        stage = "execute"
        code = (
            "from openpyxl import load_workbook; "
            "from openpyxl.chart import BarChart, Reference; "
            f"w=load_workbook({path!r}); s=w.create_sheet('Preflight'); "
            "s.append(['value']); s.append([7]); c=BarChart(); "
            "c.add_data(Reference(s,min_col=1,min_row=1,max_row=2),titles_from_data=True); "
            f"s.add_chart(c,'C1'); w.save({path!r})"
        )
        result = asyncio.run(
            asyncio.wait_for(
                terminal.execute("python3 -c " + shlex.quote(code), 10), 15
            )
        )
        if result.get("exit_code") != 0 or result.get("error"):
            raise ValueError
        stage = "publish"
        publication = asyncio.run(
            asyncio.wait_for(terminal.publish(path, "edited_in_place", 10), 15)
        )
        url = publication.get("download_url")
        if not isinstance(url, str) or not url.endswith("/preflight.xlsx"):
            raise ValueError
        stage = "download"
        artifacts = download_files([url])
        stage = "reopen"
        downloaded = load_workbook(io.BytesIO(base64.b64decode(artifacts[0]["bytes"])))
        if (
            downloaded.sheetnames != ["Sheet", "Preflight"]
            or downloaded["Sheet"]["B1"].value != 7
            or downloaded["Preflight"]["A2"].value != 7
            or len(getattr(downloaded["Preflight"], "_charts")) != 1
        ):
            raise ValueError
        downloaded.close()
    except Exception:
        error = "m24_dependency_" + stage
    finally:
        if terminal is not None:
            # The UUID directory is the only removal target, even after partial init.
            code = (
                "import shutil; from pathlib import Path; p=Path("
                + repr(terminal.root)
                + "); shutil.rmtree(p) if p.exists() else None; assert not p.exists()"
            )
            try:
                result = asyncio.run(
                    asyncio.wait_for(
                        terminal.execute("python3 -c " + shlex.quote(code), 5), 10
                    )
                )
                if result.get("exit_code") != 0 or result.get("error"):
                    raise ValueError
            except Exception:
                error = error + ":cleanup" if error else "m24_dependency_cleanup"
    if error:
        raise DependencyError(error) from None


def docker(*args: str) -> str:
    result = subprocess.run(
        ["docker", *args], capture_output=True, text=True, timeout=30
    )
    if result.returncode:
        raise DependencyError("m24_dependency_docker_" + args[0])
    return result.stdout.strip()


def resources(terminal_name: str, ui_name: str) -> list[tuple[str, str]]:
    ports = atlas_ports()
    return [
        ("container", ports.name(terminal_name)),
        ("container", ports.name(ui_name)),
        ("volume", ports.name("atlas-terminal-files")),
        ("network", ports.name(terminal_stack.NETWORK)),
    ]


def exists(kind: str, name: str) -> bool:
    if kind == "container":
        return bool(
            docker(
                "ps", "-a", "--filter", "name=^/" + name + "$", "--format", "{{.Names}}"
            )
        )
    return (
        name
        in docker(
            kind, "ls", "--filter", "name=" + name, "--format", "{{.Name}}"
        ).splitlines()
    )


@contextmanager
def file_dependencies(attempt: Path) -> Iterator[None]:
    """Reject shared namespaces, then start the existing terminal/UI stack."""
    settings = {
        "ATLAS_PORT_BASE": os.environ.get("ATLAS_PORT_BASE", "24000"),
        "OPEN_TERMINAL_API_KEY": secrets.token_urlsafe(32),
        "ATLAS_TERMINAL_ENABLED": "1",
        "ATLAS_STABLE": "0",
        "ATLAS_CODE_MODE": "0",
        "ATLAS_PROJECT_MODE": "0",
        "ATLAS_PROJECT_DSN": "",
        "ATLAS_WEBUI_URL": "",
        "ATLAS_TTS": "0",
        "ATLAS_STT": "0",
    }
    with patch.dict(os.environ, settings):
        # UI startup uses setdefault for its private URL; do not inherit another UI.
        os.environ.pop("ATLAS_WEBUI_URL", None)
        ports = atlas_ports()
        if ports.base < 10000:
            raise DependencyError("m24_dependency_dedicated_port_offset_required")
        for service in ("terminal", "chat_ui", "adapter"):
            try:
                with socket.socket() as listener:
                    listener.bind(("127.0.0.1", ports.port(service)))
            except OSError:
                raise DependencyError("m24_dependency_port_busy_" + service) from None
        owner = hashlib.sha256(str(Path.cwd().resolve()).encode()).hexdigest()[:12]
        terminal_name, ui_name = (
            "atlas-refresh-terminal-" + owner,
            "atlas-refresh-ui-" + owner,
        )
        owned = resources(terminal_name, ui_name)
        if any(exists(kind, name) for kind, name in owned):
            raise DependencyError("m24_dependency_resource_already_exists")
        (attempt / "resources.json").write_text(json.dumps(owned))
        try:
            with ExitStack() as stack:
                stage = "terminal_start"
                try:
                    stack.enter_context(
                        terminal_stack.terminal(
                            terminal_name, "127.0.0.1:" + str(ports.port("terminal"))
                        )
                    )
                    stage = "ui_start"
                    stack.enter_context(terminal_stack.user_interface(ui_name, False))
                except Exception:
                    raise DependencyError("m24_dependency_" + stage) from None
                preflight()
                yield
        finally:
            # All names were absent before startup. Include partial-start failures.
            errors = []
            for kind, name in owned:
                try:
                    if exists(kind, name):
                        docker(
                            kind, "rm", *(["-f"] if kind == "container" else []), name
                        )
                except Exception:
                    errors.append(kind)
            if errors:
                raise DependencyError(
                    "m24_dependency_resource_cleanup_" + "_".join(errors)
                )
